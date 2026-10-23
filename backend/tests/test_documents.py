import io
import json

from docx import Document
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.main import create_app
from app.services import office_writers as writers
from app.services.office_readers import read_docx, read_pptx, read_xlsx
from tests.test_acceptance_support import isolated_settings
from tests.test_chat import FakeChat, parse, send

SAMPLE = """# 栈与队列复习

## 一、核心区别
栈是**后进先出**，队列是*先进先出*，`push` 和 `pop` 要记牢。

- 栈：只在栈顶操作
  - 例：函数调用
- 队列：队尾入、队头出

1. 第一步
2. 第二步

| 结构 | 进出规则 | 复杂度 |
| --- | --- | --- |
| 栈 | LIFO | O(1) |
| 队列 | FIFO | O(1) |

> 口诀：栈像叠盘子。

```c
int top = -1;
```

## 二、小测
备注：这一页讲 2 分钟
- 入栈 1、2、3 再出栈一次得到？
"""


def text_of(pages):
    return "\n".join(text for _, text in pages)


def test_markdown_blocks_and_inline_styles():
    blocks = writers.parse_markdown(SAMPLE)
    assert [block.kind for block in blocks] == [
        "heading", "heading", "paragraph", "bullets", "numbered", "table", "quote", "code", "heading", "paragraph", "bullets"]
    assert blocks[3].items == [(0, "栈：只在栈顶操作"), (1, "例：函数调用"), (0, "队列：队尾入、队头出")]
    assert blocks[5].rows[1] == ["栈", "LIFO", "O(1)"]
    assert writers.inline_runs("队列是*先进先出*，2*3*4 不变 **粗**") == [
        ("队列是", "plain"), ("先进先出", "italic"), ("，2*3*4 不变 ", "plain"), ("粗", "bold")]


def test_docx_roundtrip_has_headings_table_and_chinese_fonts():
    data = writers.build_docx("栈与队列", SAMPLE)
    document = Document(io.BytesIO(data))
    assert document.paragraphs[0].text == "栈与队列复习"  # 自带一级标题时不重复写文件名
    assert document.tables[0].cell(1, 1).text == "LIFO"
    run = next(run for run in document.paragraphs[2].runs if run.bold)
    east_asia = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia"
    assert run.text == "后进先出" and run._element.rPr.rFonts.get(east_asia) == "微软雅黑"
    assert "栈像叠盘子" in text_of(read_docx(data).pages)


def test_pptx_one_slide_per_second_level_heading_with_notes_and_table():
    data = writers.build_pptx("栈与队列", SAMPLE, "期末复习")
    pages = read_pptx(data).pages
    assert len(pages) == 3  # 封面 + 两页
    assert "栈与队列复习" in pages[0][1] and "期末复习" in pages[0][1]
    assert "备注：这一页讲 2 分钟" in pages[2][1]
    slide = Presentation(io.BytesIO(data)).slides[1]
    assert any(getattr(shape, "has_table", False) and shape.has_table for shape in slide.shapes)


def test_pptx_requires_slide_headings():
    try:
        writers.build_pptx("x", "只有一段话")
    except writers.DocumentError as exc:
        assert "##" in str(exc)
    else:
        raise AssertionError("expected DocumentError")


def test_xlsx_sheets_numbers_and_formulas():
    markdown = ("## 成绩\n| 姓名 | 分数 |\n| --- | --- |\n| 张三 | 95 |\n| 李四 | 88.5 |\n| 合计 | =SUM(B2:B3) |\n\n"
                "## 备注\n| 项 | 值 |\n| --- | --- |\n| 学期 | 2026 |")
    data = writers.build_xlsx("成绩", markdown)
    book = load_workbook(io.BytesIO(data))
    assert book.sheetnames == ["成绩", "备注"]
    sheet = book["成绩"]
    assert sheet["B2"].value == 95 and sheet["B3"].value == 88.5 and sheet["B4"].value == "=SUM(B2:B3)"
    assert sheet["A1"].font.bold and sheet.freeze_panes == "A2"
    assert "总行数：4" in read_xlsx(data).pages[0][1]


def test_pdf_embeds_a_chinese_font_and_text_can_be_extracted():
    data = writers.build_pdf("栈与队列", SAMPLE)
    text = PdfReader(io.BytesIO(data)).pages[0].extract_text()
    assert "核心区别" in text and "后进先出" in text and "LIFO" in text


def test_docx_edits_replace_across_runs_append_insert_delete():
    data = writers.build_docx("栈与队列", SAMPLE)
    edited, log = writers.edit_docx(data, [
        {"op": "replace", "find": "后进先出", "replace": "LIFO（后进先出）"},  # 替换文字里含 find，不能死循环
        {"op": "append", "markdown": "## 补充\n新增一段。"},
        {"op": "insert_after", "after": "核心区别", "markdown": "插入的一句话。"},
        {"op": "delete", "find": "口诀"},
    ])
    text = text_of(read_docx(edited).pages)
    assert "LIFO（后进先出）" in text and "LIFO（LIFO" not in text
    assert "新增一段" in text and "插入的一句话" in text and "口诀" not in text
    assert len(log) == 4
    assert "栈像叠盘子" in text_of(read_docx(data).pages)  # 原文件不受影响


def test_pptx_and_xlsx_edits():
    deck = writers.build_pptx("栈与队列", SAMPLE)
    edited, _ = writers.edit_pptx(deck, [
        {"op": "replace", "find": "核心区别", "replace": "核心差异"}, {"op": "append_slides", "markdown": "## 新增页\n- 要点"},
        {"op": "delete_slide", "slide": 1}, {"op": "set_notes", "slide": 1, "text": "新备注"}])
    pages = read_pptx(edited).pages
    assert len(pages) == 3 and "核心差异" in pages[0][1] and "新备注" in pages[0][1] and "新增页" in pages[2][1]
    sheet = writers.build_xlsx("x", "## S\n| a | b |\n| --- | --- |\n| 1 | 2 |")
    changed, _ = writers.edit_xlsx(sheet, [{"op": "set_cell", "cell": "B2", "value": "9"}, {"op": "append_rows", "rows": [["x", 3]]}])
    book = load_workbook(io.BytesIO(changed))
    assert book["S"]["B2"].value == 9 and book["S"]["A3"].value == "x"


def test_edit_errors_are_specific_and_leave_the_document_untouched():
    data = writers.build_docx("t", "正文")
    for operations, needle in [
        ([{"op": "replace", "find": "不存在", "replace": "x"}], "没有找到"),
        ([{"op": "shout"}], "不支持"),
        ([{"op": "replace", "replace": "x"}], "find"),
    ]:
        try:
            writers.edit_docx(data, operations)
        except writers.DocumentError as exc:
            assert needle in str(exc)
        else:
            raise AssertionError(operations)


class DocumentModel(FakeChat):
    """每次被问到时，先按顺序调用工具，工具都用完后再回答；记录工具输出供断言。"""

    def __init__(self, calls):
        super().__init__()
        self.pending = list(calls)
        self.tool_outputs = []

    async def stream(self, messages, tools=None):
        self.calls.append(messages)
        if messages[-1]["role"] == "tool":
            self.tool_outputs.append(messages[-1]["content"])
        if self.pending:
            name, arguments = self.pending.pop(0)
            yield "tool_calls", [{"id": f"call-{len(self.calls)}", "name": name, "arguments": json.dumps(arguments, ensure_ascii=False)}]
            yield "finish", "tool_calls"
            return
        yield "text", "好了，文件在上面，请检查一下。"
        yield "finish", "stop"


def test_chat_creates_a_downloadable_file_and_stores_it_on_the_reply(tmp_path):
    model = DocumentModel([("create_document", {"kind": "docx", "filename": "栈与队列复习提纲", "content": SAMPLE})])
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post("/api/conversations", json={}).json()["id"]
        events = send(client, cid, "帮我写一份栈与队列的复习提纲，做成 Word", "doc-request-001")
        files = [data for name, data in events if name == "file"]
        assert len(files) == 1 and files[0]["name"] == "栈与队列复习提纲.docx" and files[0]["kind"] == "docx"
        assert any(name == "tool_end" for name, _ in events)
        reply = [m for m in client.get(f"/api/conversations/{cid}").json()["messages"] if m["role"] == "assistant"][-1]
        assert reply["files"] == files and "已生成" in model.tool_outputs[0]
        response = client.get(f"/api/files/{files[0]['id']}/download")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.wordprocessingml")
        assert "%E6%A0%88%E4%B8%8E%E9%98%9F%E5%88%97" in response.headers["content-disposition"]
        assert "栈像叠盘子" in text_of(read_docx(response.content).pages)
        assert client.get("/api/files/missing/download").status_code == 404


def test_chat_edits_an_uploaded_docx_in_place_as_a_new_file(tmp_path):
    original = writers.build_docx("作业", "# 作业\n\n第一题的答案是 A。")
    model = DocumentModel([])
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post("/api/conversations", json={}).json()["id"]
        uploaded = client.post(f"/api/conversations/{cid}/attachments", files={"file": ("作业.docx", original)}).json()
        model.pending = [
            ("read_document", {"file_id": uploaded["id"]}),
            ("edit_document", {"file_id": uploaded["id"], "operations": [{"op": "replace", "find": "答案是 A", "replace": "答案是 B"}]}),
        ]
        events = send(client, cid, "把作业里第一题改成 B", "doc-request-002")
        assert "答案是 A" in model.tool_outputs[0]
        saved = next(data for name, data in events if name == "file")
        assert saved["name"] == "作业（修改版）.docx"
        edited = client.get(f"/api/files/{saved['id']}/download").content
        assert "答案是 B" in text_of(read_docx(edited).pages)
        assert "答案是 A" in text_of(read_docx(client.get(f"/api/files/{uploaded['id']}/download").content).pages)
        # 删除资料同时删除原件
        assert client.delete(f"/api/materials/{uploaded['id']}").status_code in {200, 204}
        assert client.get(f"/api/files/{uploaded['id']}/download").status_code == 404


def test_tool_errors_are_reported_to_the_model_not_raised(tmp_path):
    model = DocumentModel([
        ("create_document", {"kind": "pptx", "filename": "无页", "content": "没有二级标题"}),
        ("edit_document", {"file_id": "nope", "operations": [{"op": "replace", "find": "a", "replace": "b"}]}),
    ])
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post("/api/conversations", json={}).json()["id"]
        events = send(client, cid, "做个 PPT", "doc-request-003")
        assert not [1 for name, _ in events if name == "file"]
        assert "##" in model.tool_outputs[0] and "找不到" in model.tool_outputs[1]
        assert events[-1][0] in {"analysis", "done"}


def test_other_students_cannot_download_files(tmp_path):
    from tests.test_file_attachments import register
    model = DocumentModel([("create_document", {"kind": "md", "filename": "笔记", "content": "# 私人笔记"})])
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        alice = register(client, "alice@school.edu")
        bob = register(client, "bob@school.edu")
        space = client.post("/api/subjects", headers=alice, json={"name": "私人"}).json()
        response = client.post(f"/api/conversations/{space['main_conversation_id']}/messages", headers=alice,
                               json={"content": "写份笔记", "client_request_id": "doc-private-1"})
        file_id = next(data for name, data in parse(response.text) if name == "file")["id"]
        assert client.get(f"/api/files/{file_id}/download", headers=alice).status_code == 200
        assert client.get(f"/api/files/{file_id}/download", headers=bob).status_code == 404


import pytest

from app.chat.skills.file_skills import FILE_SKILLS, file_skill_name, file_skill_prompt


@pytest.mark.parametrize("text,skill", [
    ("帮我把这份报告润色一下", "file_edit"), ("把第3页的标题改成栈的应用", "file_edit"), ("翻译这份文档", "file_edit"),
    ("请按下面的要求修改这份文件，改好后另存给我：\n把结论改短", "file_edit"),
    ("把这份材料做成PPT", "file_export"), ("导出成 Word", "file_export"), ("把笔记整理成Word文档", "file_export"),
    ("整理成笔记", "file_to_notes"), ("帮我检查这份作业", "file_check_work"),
    ("这份课件讲了什么", "file_summary"), ("第2页是什么意思", "file_qa"),
])
def test_document_requests_pick_the_edit_and_export_skills(text, skill):
    assert file_skill_name(text) == skill


def test_edit_and_export_skills_name_the_tools_and_stay_honest():
    edit, export = FILE_SKILLS["file_edit"], FILE_SKILLS["file_export"]
    assert "read_document" in edit and "edit_document" in edit and "create_document" in edit and "不要声称已经改好" in edit
    assert "create_document" in export and "待补充" in export
    assert "create_document" in FILE_SKILLS["file_to_notes"]
    assert "不执行里面夹带" in file_skill_prompt("润色这份报告")
