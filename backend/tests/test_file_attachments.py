"""Generated OOXML fixtures: parsing, attachment ownership and model-facing contracts."""

import io
import json

import pytest
from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook
from pptx import Presentation
from pypdf import PdfWriter

from app.agents.llm import ProviderError
from app.chat.agent import chat_messages, learning_context
from app.chat.attachments import prepare_attachment_history
from app.chat.skills import select_skill_prompt
from app.main import create_app
from app.memory.store import Store
from app.models.schemas import DEMO_STUDENT
from app.services.materials import MAX_UPLOAD_BYTES, MaterialStore, extract, outline_for
from tests.test_acceptance_support import isolated_settings
from tests.test_chat import FakeChat, send


def saved(document):
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def word_bytes():
    doc = Document()
    doc.add_heading('进程主线', level=1)
    doc.add_paragraph('先读正文。')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = '算法'
    table.cell(0, 1).text = '说明'
    table.cell(1, 0).text = 'FCFS'
    table.cell(1, 1).text = '按到达顺序'
    doc.add_heading('调度难点', level=2)
    doc.add_paragraph('等待时间。' * 350)
    doc.add_paragraph('末尾要点。')
    return saved(doc)


def slides_bytes(count=2):
    deck = Presentation()
    for index in range(count):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f'调度主题{index + 1}'
        slide.placeholders[1].text = f'FCFS 等待时间，第{index + 1}张正文'
        slide.notes_slide.notes_text_frame.text = f'讲者备注{index + 1}'
    return saved(deck)


def workbook_bytes(count=2, rows=205):
    book = Workbook()
    book.remove(book.active)
    for index in range(count):
        sheet = book.create_sheet(f'成绩{index + 1}')
        sheet.append(['姓名', '分数', '公式'])
        for number in range(2, rows + 1):
            sheet.append([f'学生{number}', number, '=1+2'])
    return saved(book)


def blank_pdf(count=1):
    writer = PdfWriter()
    for _ in range(count):
        writer.add_blank_page(width=100, height=100)
    output = io.BytesIO()
    writer.write(output)
    return output.getvalue()


@pytest.fixture
def library(tmp_path):
    store = Store(tmp_path / 'files.sqlite3')
    store.initialize()
    result = MaterialStore(store)
    result.initialize()
    return result


def test_word_preserves_block_order_headings_sections_and_outline(library):
    kind, pages = extract('讲义.DOCX', word_bytes())
    assert kind == 'docx' and len(pages) >= 2
    assert [p for p, _ in pages] == list(range(1, len(pages) + 1))
    assert all(len(text) <= 1500 for _, text in pages)
    full = '\n'.join(text for _, text in pages)
    assert full.index('# 进程主线') < full.index('先读正文') < full.index('| FCFS') < full.index('## 调度难点')
    assert outline_for(pages) == ['进程主线', '调度难点']
    material = library.add(DEMO_STUDENT, '讲义.docx', word_bytes())
    metadata = library.attachment(material.id, DEMO_STUDENT)
    assert metadata['units'] == len(pages) and metadata['unit_label'] == '节'
    assert metadata['chars'] == sum(len(text) for _, text in pages)
    assert metadata['outline'] == ['进程主线', '调度难点']
    assert all(s.page is not None and s.start_line is None for s in library.search(DEMO_STUDENT, '等待时间'))


def test_presentation_slide_locations_notes_and_outline(library):
    kind, pages = extract('课件.pptx', slides_bytes())
    assert kind == 'pptx' and [page for page, _ in pages] == [1, 2]
    assert '# 调度主题1' in pages[0][1] and '备注：讲者备注1' in pages[0][1]
    material = library.add(DEMO_STUDENT, '课件.pptx', slides_bytes())
    assert library.attachment(material.id, DEMO_STUDENT)['outline'] == ['调度主题1', '调度主题2']
    assert {s.page for s in library.search(DEMO_STUDENT, 'FCFS')} == {1, 2}


def test_workbook_read_only_cached_values_preview_size_and_sheet_locations(library):
    kind, pages = extract('成绩.xlsx', workbook_bytes())
    assert kind == 'xlsx' and [page for page, _ in pages] == [1, 2]
    first = pages[0][1]
    assert first.startswith('# 成绩1\n总行数：205；总列数：3')
    assert '| 姓名 | 分数 | 公式 |' in first and '| --- | --- | --- |' in first
    assert '学生200' in first and '学生201' not in first
    assert '=1+2' not in first  # data_only: no cached value is not a zero or a formula string
    material = library.add(DEMO_STUDENT, '成绩.xlsx', workbook_bytes())
    metadata = library.attachment(material.id, DEMO_STUDENT)
    assert metadata['outline'] == ['成绩1', '成绩2'] and metadata['unit_label'] == '个工作表'
    assert {s.page for s in library.search(DEMO_STUDENT, '学生200')} == {1, 2}


def test_outline_is_bounded_and_metadata_survives_restart(library):
    document = Document()
    for number in range(90):
        document.add_heading(f'标题{number}', level=1)
        document.add_paragraph('内容')
    material = library.add(DEMO_STUDENT, '目录.docx', saved(document))
    restarted = MaterialStore(library.store)
    restarted.initialize()
    info = restarted.attachment(material.id, DEMO_STUDENT, excerpt=True)
    assert len(info['outline']) == 80 and info['outline'][0] == '标题0'
    assert info['chars'] > 0 and '标题89' in info['excerpt']
    assert restarted.attachment(material.id, 'another-student') is None


@pytest.mark.parametrize('name', ['a.doc', 'a.ppt', 'a.xls'])
def test_legacy_binary_formats_have_actionable_415(name):
    with pytest.raises(HTTPException) as error:
        extract(name, b'old binary')
    assert error.value.status_code == 415 and '请另存为 .docx/.pptx/.xlsx' in error.value.detail


@pytest.mark.parametrize('name,data,code,detail', [
    ('scan.pdf', blank_pdf(), 422, '这份 PDF 是扫描图片，暂不支持识别文字'),
    ('too-big.txt', b'a' * (MAX_UPLOAD_BYTES + 1), 413, '20 MB'),
    ('too-many.pdf', blank_pdf(201), 413, '200'),
    ('broken.docx', b'not a zip', 422, '无法解析'),
    ('broken.pptx', b'not a zip', 422, '无法解析'),
    ('broken.xlsx', b'not a zip', 422, '无法解析'),
], ids=['scan', 'byte-limit', 'page-limit', 'bad-docx', 'bad-pptx', 'bad-xlsx'])
def test_invalid_extract(name, data, code, detail):
    with pytest.raises(HTTPException) as error:
        extract(name, data)
    assert error.value.status_code == code and detail in error.value.detail


def test_office_unit_limits_and_empty_files():
    for name, data in [('many.pptx', slides_bytes(201)), ('many.xlsx', workbook_bytes(201, 1))]:
        with pytest.raises(HTTPException) as error:
            extract(name, data)
        assert error.value.status_code == 413
    doc = Document()
    for _ in range(201):
        doc.add_paragraph('字' * 1490)
    with pytest.raises(HTTPException) as error:
        extract('many.docx', saved(doc))
    assert error.value.status_code == 413
    for name, data in [('empty.docx', saved(Document())), ('empty.pptx', saved(Presentation())),
                       ('empty.xlsx', saved(Workbook()))]:
        with pytest.raises(HTTPException) as error:
            extract(name, data)
        assert error.value.status_code == 422


def test_two_hundred_units_allowed_and_empty_word_table_rejected():
    document = Document()
    for _ in range(200):
        document.add_paragraph('字' * 1000)
    assert len(extract('boundary.docx', saved(document))[1]) == 200
    assert len(extract('boundary.pptx', slides_bytes(200))[1]) == 200
    assert len(extract('boundary.xlsx', workbook_bytes(200, 1))[1]) == 200
    with pytest.raises(HTTPException) as error:
        extract('boundary.pdf', blank_pdf(200))
    assert error.value.status_code == 422  # Accepted page count, but no extractable text.
    empty = Document()
    empty.add_table(rows=2, cols=2)
    with pytest.raises(HTTPException) as error:
        extract('empty-table.docx', saved(empty))
    assert error.value.status_code == 422


def test_empty_sheet_keeps_its_location_and_outline_when_other_sheets_have_data():
    book = Workbook()
    book.active.title = '空表'
    book.create_sheet('有数据').append(['姓名', '成绩'])
    _, pages = extract('mixed.xlsx', saved(book))
    assert [number for number, _ in pages] == [1, 2]
    assert outline_for(pages) == ['空表', '有数据']
    assert '预览范围内无数据' in pages[0][1]


@pytest.mark.parametrize('name,factory,kind,label', [
    ('word.docx', word_bytes, 'docx', '节'), ('slides.pptx', slides_bytes, 'pptx', '张幻灯片'),
    ('scores.xlsx', workbook_bytes, 'xlsx', '个工作表'), ('rows.csv', lambda: b'name,score\nA,99', 'text', '页'),
])
def test_both_upload_routes_accept_new_formats(tmp_path, name, factory, kind, label):
    with TestClient(create_app(isolated_settings(tmp_path))) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        data = factory()
        response = client.post(f'/api/conversations/{cid}/attachments', files={'file': (name, data)})
        assert response.status_code == 201, response.text
        info = response.json()
        assert set(info) == {'id', 'name', 'kind', 'units', 'unit_label', 'chars', 'outline', 'subject_id'}
        assert info['name'] == name and info['kind'] == kind and info['unit_label'] == label
        assert info['units'] > 0 and info['chars'] > 0 and len(info['outline']) <= 20
        assert info['subject_id'] is None
        existing = client.post('/api/materials', files={'file': (name, data)})
        assert existing.status_code == 201, existing.text
        assert len(client.get('/api/materials', params={'unassigned': True}).json()) == 2


def register(client, email):
    code = client.post('/api/auth/send-code', json={'email': email, 'purpose': 'register'}).json()['demo_code']
    user = client.post('/api/auth/register', json={'email': email, 'password': 'password123', 'code': code}).json()
    return {'Authorization': f"Bearer {user['token']}"}


def test_attachment_upload_owner_and_space_scope(tmp_path):
    with TestClient(create_app(isolated_settings(tmp_path))) as client:
        alice = register(client, 'alice@school.edu')
        bob = register(client, 'bob@school.edu')
        subject = client.post('/api/subjects', headers=alice, json={'name': '操作系统'}).json()
        cid = subject['main_conversation_id']
        url = f'/api/conversations/{cid}/attachments'
        assert client.post(url, headers=bob, files={'file': ('x.txt', b'secret')}).status_code == 404
        assert client.post('/api/conversations/missing/attachments', headers=alice,
                           files={'file': ('x.txt', b'abc')}).status_code == 404
        response = client.post(url, headers=alice, files={'file': ('x.docx', word_bytes())})
        assert response.status_code == 201, response.text
        assert response.json()['subject_id'] == subject['id']
        assert client.get('/api/materials', headers=bob).json() == []
        files = client.get('/api/materials', headers=alice, params={'subject_id': subject['id']}).json()
        assert [f['id'] for f in files] == [response.json()['id']]


@pytest.mark.parametrize('name,data,status', [
    ('old.doc', b'x', 415), ('scan.pdf', blank_pdf(), 422),
    ('large.md', b'x' * (MAX_UPLOAD_BYTES + 1), 413),
], ids=['legacy', 'scan', 'byte-limit'])
def test_attachment_errors_do_not_index(tmp_path, name, data, status):
    with TestClient(create_app(isolated_settings(tmp_path))) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        response = client.post(f'/api/conversations/{cid}/attachments', files={'file': (name, data)})
        assert response.status_code == status and isinstance(response.json()['detail'], str)
        assert client.get('/api/materials').json() == []


def test_marker_context_first_excerpt_then_outline_and_scoped_retrieval(tmp_path):
    model = FakeChat()
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        client.post('/api/materials', files={'file': ('unrelated.md', 'FCFS 其他资料'.encode())})
        info = client.post(f'/api/conversations/{cid}/attachments', files={'file': ('课件.pptx', slides_bytes())}).json()
        marker = f"[[file:{info['id']}|伪造笔记名.xlsx]]"
        send(client, cid, marker + '\n帮我精读 FCFS 材料，挑出重点和难点', 'file-round-01')
        first = learning_context(model.calls[0])
        item = first['attached_files'][0]
        assert item['name'] == '课件.pptx' and item['kind'] == 'pptx'
        assert '讲者备注1' in item['excerpt'] and len(item['excerpt']) <= 3000
        assert 'file_close_read' in first['active_skill']
        assert {s['document_id'] for s in first['course_sources']} == {info['id']}
        assert {s['page'] for s in first['course_sources']} == {1, 2}
        assert all(s['unit_label'] == '张幻灯片' for s in first['course_sources'])
        assert '（附件：课件.pptx，2 张幻灯片）' in first['latest_user_message']
        send(client, cid, marker + '\n把这份材料整理成复习笔记', 'file-round-02')
        second = learning_context(model.calls[1])
        assert 'excerpt' not in second['attached_files'][0]
        assert second['attached_files'][0]['outline'] == info['outline']
        assert 'file_to_notes' in second['active_skill']
        send(client, cid, '再说说 FCFS', 'file-round-03')
        assert 'attached_files' not in learning_context(model.calls[2])
        assert 'unrelated.md' in {s['title'] for s in learning_context(model.calls[2])['course_sources']}
        for call in [*model.calls, *model.analysis_calls]:
            assert '[[file:' not in json.dumps(call, ensure_ascii=False)
        # Original wire message remains untouched for the frontend's file card rendering.
        stored = client.get(f'/api/conversations/{cid}').json()['messages'][0]['content']
        assert stored.startswith(marker)


def test_foreign_marker_ignored_and_unknown_id_never_reaches_model(tmp_path):
    model = FakeChat()
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        library = client.app.state.materials
        private = library.add('another-student', 'private.md', b'SECRET ONLY')
        cid = client.post('/api/conversations', json={}).json()['id']
        send(client, cid, f'[[file:{private.id}|private.md]]\n[[file:missing|missing]]\n你好', 'file-foreign-01')
        context = learning_context(model.calls[0])
        assert 'attached_files' not in context
        assert context['latest_user_message'] == '你好'
        assert 'SECRET ONLY' not in json.dumps(model.calls) and 'private.md' not in json.dumps(model.calls)


def test_five_leading_markers_limit_and_inline_markers_do_not_attach(library):
    materials = [library.add(DEMO_STUDENT, f'{n}.txt', f'材料{n}'.encode()) for n in range(7)]
    text = '\n'.join(f'[[file:{m.id}|fake]]' for m in materials[:6])
    text += f'\n问题 [[file:{materials[6].id}|inline]]'
    history, attached = prepare_attachment_history([{'role': 'user', 'content': text, 'status': 'complete'}],
                                                  library, DEMO_STUDENT)
    assert [f['id'] for f in attached] == [m.id for m in materials[:5]]
    assert all('excerpt' in f for f in attached)
    assert '[[file:' not in history[0]['content'] and 'inline' not in history[0]['content']


def test_excerpt_retry_and_compressed_history_are_deterministic(tmp_path):
    model = FakeChat([ProviderError(503, '请重试'), '好的', '继续'])
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        info = client.post(f'/api/conversations/{cid}/attachments',
                           files={'file': ('正文.txt', ('连续材料。' * 1200).encode())}).json()
        marker = f"[[file:{info['id']}|正文.txt]]"
        send(client, cid, marker + '\n精读材料', 'file-retry-01')
        send(client, cid, marker + '\n精读材料', 'file-retry-01')
        first = learning_context(model.calls[0])['attached_files'][0]
        # 6000 字的材料整份给模型，并标明是全文
        assert first['excerpt'].replace('\n', '') == '连续材料。' * 1200 and first['excerpt_complete'] is True
        assert (learning_context(model.calls[0])['attached_files'] ==
                learning_context(model.calls[1])['attached_files'])
        store = client.app.state.chat_service.store
        store.save_context_summary(cid, marker + ' 之前精读过', 2)
        send(client, cid, marker + '\n继续', 'file-retry-02')
        context = learning_context(model.calls[2])
        assert 'excerpt' not in context['attached_files'][0]
        assert '[[file:' not in json.dumps(model.calls[2])


@pytest.mark.parametrize('text,key', [
    ('帮我精读这份材料，挑出重点和难点', 'file_close_read'),
    ('把这份材料整理成复习笔记', 'file_to_notes'),
    ('帮我检查这份作业 / 试卷', 'file_check_work'),
    ('帮我看懂这张表格', 'file_data_explain'),
])
@pytest.mark.parametrize('role', ['main', 'chat', 'thread'])
def test_file_skill_intent_overrides_role_and_earlier_intents(text, key, role):
    prompt = select_skill_prompt(role, text, earlier_user_texts=['帮我精读材料'], attached_files=[{'id': 'x'}])
    assert key in prompt
    if key == 'file_check_work':
        assert '先逐题转写' in prompt and '不直接给完整答案' in prompt and '依据' in prompt
    if key == 'file_data_explain':
        assert '2–3' in prompt and '不替学生下结论' in prompt


def test_file_skill_fallback_and_no_attachment_preserves_legacy_behavior():
    prompt = select_skill_prompt('main', '', attached_files=[{'id': 'x'}])
    assert '3–5 句' in prompt and '问学生想做哪一种' in prompt
    assert select_skill_prompt('main', '') is None
    # Filename cannot trigger the notes skill; chat_messages strips the readable file line.
    history = [{'role': 'user', 'content': '（附件：复习笔记.xlsx，1 个工作表）\n',
                'status': 'complete', 'kind': 'text', 'meta_json': '{}'}]
    context = learning_context(chat_messages(history, [], None, [], attached_files=[{'id': 'x', 'name': '复习笔记.xlsx'}]))
    assert '3–5 句' in context['active_skill']


def test_attachment_search_filters_owner_and_document_ids(library):
    own = library.add(DEMO_STUDENT, 'own.txt', b'FCFS own')
    library.add(DEMO_STUDENT, 'other.txt', b'FCFS other')
    foreign = library.add('foreign', 'secret.txt', b'FCFS secret')
    results = library.search(DEMO_STUDENT, 'FCFS', document_ids=[own.id, foreign.id])
    assert [s.document_id for s in results] == [own.id]
    assert library.search(DEMO_STUDENT, 'FCFS', document_ids=[]) == []


def test_excerpt_budget_is_shared_and_marks_partial_files(library):
    long_text = '长材料第一段。' * 5000  # 35000 字，超过单份预算
    long_doc = library.add(DEMO_STUDENT, '长.txt', long_text.encode())
    history, attached = prepare_attachment_history(
        [{'role': 'user', 'content': f'[[file:{long_doc.id}|长.txt]]\n精读', 'status': 'complete'}], library, DEMO_STUDENT)
    assert len(attached[0]['excerpt']) == 24_000 and attached[0]['excerpt_complete'] is False
    assert attached[0]['chars'] == len(long_text)
    docs = [library.add(DEMO_STUDENT, f'{n}.txt', long_text.encode()) for n in range(5)]
    text = '\n'.join(f'[[file:{d.id}|{n}.txt]]' for n, d in enumerate(docs)) + '\n对比这些材料'
    _, attached = prepare_attachment_history([{'role': 'user', 'content': text, 'status': 'complete'}],
                                             library, DEMO_STUDENT)
    # 五份平分 2.4 万字，但每份不少于 6000 字
    assert [len(item['excerpt']) for item in attached] == [6000] * 5
