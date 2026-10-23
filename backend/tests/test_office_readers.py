"""文件读取：各格式（含旧版 ppt / doc / xls）解析结果、读取说明、技能选择，以及到模型上下文的完整链路。

tests/fixtures/office 里的 .ppt / .doc / .xls / .docx / .pptx / .xlsx / .pdf 都是用真实的 Office 软件生成的，
同一份内容分别存成新旧两种格式，用来核对旧格式解析和新格式一致。
"""

import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches
from pypdf import PdfReader, PdfWriter

from app.chat.agent import learning_context
from app.chat.attachments import prepare_attachment_history
from app.chat.skills.file_skills import file_skill_name, file_skill_prompt
from app.main import create_app
from app.memory.store import Store
from app.models.schemas import DEMO_STUDENT
from app.services import office_readers as readers
from app.services.materials import MaterialStore, extract, extract_document
from tests.test_acceptance_support import isolated_settings
from tests.test_chat import FakeChat, send

FIXTURES = Path(__file__).parent / 'fixtures' / 'office'


def fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


def joined(reading) -> str:
    return '\n'.join(text for _, text in reading.pages)


@pytest.fixture
def library(tmp_path):
    store = Store(tmp_path / 'files.sqlite3')
    store.initialize()
    materials = MaterialStore(store)
    materials.initialize()
    return materials


# ---------------------------------------------------------------- PowerPoint

DECK_FACTS = ['数据结构 第3章 栈和队列', '只允许在栈顶插入和删除', '后进先出 LIFO', '空栈 pop 会下溢',
              '误区：栈和队列一样是先进先出', '组合里的文字A', '组合里的文字B', '栈顶', '队头',
              '备用页：拓展阅读', '双端队列与单调栈', '讲到这页时让学生先自己画一遍']


@pytest.mark.parametrize('name', ['deck.pptx', 'deck.ppt'])
def test_deck_text_is_complete_in_new_and_legacy_format(name):
    reading = (readers.read_pptx if name.endswith('x') else readers.read_ppt)(fixture(name))
    text = joined(reading)
    assert [n for n, _ in reading.pages] == [1, 2, 3, 4, 5]
    for fact in DECK_FACTS:
        assert fact in text, (name, fact)
    assert reading.pages[1][1].startswith('# 栈的定义')
    assert '备注：讲到这页时让学生先自己画一遍' in reading.pages[3][1]
    # 隐藏页：文字读到，并标明是隐藏的
    assert '（这张幻灯片在放映时是隐藏的）' in reading.pages[4][1]


def test_pptx_keeps_structure_levels_tables_and_picture_alt_text():
    reading = readers.read_pptx(fixture('deck.pptx'))
    assert '- 后进先出 LIFO' in reading.pages[1][1]            # 第二级要点缩进
    assert '| 结构 | 入口 | 出口 |' in reading.pages[2][1] and '| --- | --- | --- |' in reading.pages[2][1]
    assert '| 队列 | 队尾 | 队头 |' in reading.pages[2][1]
    assert '[图片：入栈 1 2 3 的出栈示意]' in reading.pages[3][1]
    assert any('第 4 张幻灯片含图片' in note and '没有读取' in note for note in reading.notes)


def test_legacy_ppt_says_what_it_cannot_read():
    reading = readers.read_ppt(fixture('deck.ppt'))
    assert reading.notes == ['这份课件含图片，图片、图表里的文字没有读取。']
    # 没有图片的课件不应该出现「没读到」的提示
    assert readers.read_ppt(fixture('free.ppt')).notes == []


def test_notes_only_list_content_that_was_really_skipped():
    for name, reader in [('free.ppt', readers.read_ppt), ('free.pptx', readers.read_pptx),
                         ('notes_plain.docx', readers.read_docx), ('scores.xls', readers.read_xls)]:
        if name == 'notes_plain.docx':
            from docx import Document
            doc = Document()
            doc.add_paragraph('只有文字的文档')
            buffer = io.BytesIO()
            doc.save(buffer)
            data = buffer.getvalue()
        else:
            data = fixture(name)
        notes = reader(data).notes
        assert all('隐藏的工作表' in note for note in notes), (name, notes)   # 只剩真的没读的隐藏表


def test_pptx_chart_and_smartart_text_are_read():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = '成绩对比'
    data = CategoryChartData()
    data.categories = ['栈', '队列']
    data.add_series('平均分', (15, 12))
    chart = slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(2), Inches(5), Inches(3), data).chart
    chart.has_title = True
    chart.chart_title.text_frame.text = '期中平均分'
    buffer = io.BytesIO()
    deck.save(buffer)
    text = joined(readers.read_pptx(buffer.getvalue()))
    assert '[图表：期中平均分；系列：平均分；类别：栈、队列]' in text

    dgm = 'http://schemas.openxmlformats.org/drawingml/2006/diagram'
    rel = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    frame = etree.fromstring(f'<p:graphicFrame xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                             f'xmlns:dgm="{dgm}" xmlns:r="{rel}"><dgm:relIds r:dm="rId9"/></p:graphicFrame>')
    blob = (f'<dgm:dataModel xmlns:dgm="{dgm}" xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"><dgm:ptLst>'
            '<dgm:pt type="doc"><dgm:t><a:p><a:r><a:t>不该出现</a:t></a:r></a:p></dgm:t></dgm:pt>'
            '<dgm:pt><dgm:t><a:p><a:r><a:t>入栈</a:t></a:r></a:p></dgm:t></dgm:pt>'
            '<dgm:pt><dgm:t><a:p><a:r><a:t>出栈</a:t></a:r></a:p></dgm:t></dgm:pt></dgm:ptLst></dgm:dataModel>').encode()
    part = SimpleNamespace(related_part=lambda rid: SimpleNamespace(blob=blob) if rid == 'rId9' else None)
    assert readers._smartart_text(SimpleNamespace(_element=frame), part) == ['入栈', '出栈']


def test_long_deck_outline_reaches_every_slide_and_excerpt_budget_is_shared(library):
    deck = Presentation()
    for index in range(60):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f'第{index + 1}讲标题'
        slide.placeholders[1].text = '内容。' * 300
    buffer = io.BytesIO()
    deck.save(buffer)
    material = library.add(DEMO_STUDENT, '长课件.pptx', buffer.getvalue())
    info = library.attachment(material.id, DEMO_STUDENT, excerpt=True, excerpt_limit=24_000)
    assert len(info['outline']) == 60 and info['outline'][-1] == '第60讲标题'
    assert info['units'] == 60 and info['excerpt_complete'] is False and len(info['excerpt']) == 24_000


# ---------------------------------------------------------------- Word

@pytest.mark.parametrize('name', ['notes.docx', 'notes.doc'])
def test_word_text_is_complete_in_new_and_legacy_format(name):
    reading = (readers.read_docx if name.endswith('x') else readers.read_doc)(fixture(name))
    text = joined(reading)
    for fact in ['队列学习笔记', '一、定义', '先进先出 FIFO', '队头出队', '队尾入队', '(rear+1)%N == front', '链队列',
                 '无需判满', '结尾：删除到空时别忘了把 tail 置空。', '【脚注】', 'tail 指针在链队列里容易悬空',
                 '页眉：数据结构复习资料', '页脚：仅供课堂使用']:
        assert fact in text, (name, fact)
    assert 'rear+1)%N == front' in text and '|' in text            # 表格按行输出，列用 | 分隔


def test_docx_structure_headings_lists_and_table_markdown():
    text = joined(readers.read_docx(fixture('notes.docx')))
    assert '# 一、定义' in text and '- 队头出队' in text
    assert '| 实现 | 判满条件 |' in text and '| --- | --- |' in text and '| 循环数组 | (rear+1)%N == front |' in text


def test_docx_floating_textbox_and_merged_cells():
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=2, cols=3)
    table.cell(0, 0).merge(table.cell(0, 2)).text = '合并的标题'
    for index, label in enumerate(['甲', '乙', '丙']):
        table.cell(1, index).text = label
    buffer = io.BytesIO()
    doc.save(buffer)
    text = joined(readers.read_docx(buffer.getvalue()))
    assert text.count('合并的标题') == 1 and '| 甲 | 乙 | 丙 |' in text


def test_doc_reports_images_only_when_there_are_some():
    assert any('1 处图片' in note for note in readers.read_doc(fixture('notes.doc')).notes)


def test_encrypted_doc_gets_a_clear_message(monkeypatch):
    import olefile
    real = olefile.OleFileIO(io.BytesIO(fixture('notes.doc')))
    word = bytearray(real.openstream('WordDocument').read())
    word[0x0B] |= 0x01                                      # FIB.flags 的 fEncrypted 位（0x0100）

    class Locked:
        def exists(self, name):
            return real.exists(name)

        def openstream(self, name):
            return io.BytesIO(bytes(word)) if name == 'WordDocument' else real.openstream(name)

        def close(self):
            real.close()

    monkeypatch.setattr(readers, '_ole', lambda data: Locked())
    with pytest.raises(HTTPException) as error:
        readers.read_doc(b'ignored')
    assert error.value.status_code == 422 and '已加密' in error.value.detail


# ---------------------------------------------------------------- Excel

@pytest.mark.parametrize('name', ['scores.xlsx', 'scores.xls'])
def test_workbook_values_formulas_dates_in_new_and_legacy_format(name):
    reading = (readers.read_xlsx if name.endswith('x') else readers.read_xls)(fixture(name))
    first = reading.pages[0][1]
    assert '| 学号 | 姓名 | 栈 | 队列 | 哈希 | 总分 |' in first
    assert '| 2026001 | 张同学 | 18 | 15 | 12 | 45 |' in first     # 公式显示计算结果，不是 =SUM(...)
    assert '=SUM' not in first and '2026-10-24' in first         # 日期不是序列号
    assert '合并单元格：备注说明' in first
    assert '循环队列判满 | 9' in reading.pages[1][1]
    # 表格宽度就是实际用到的 6 列（7 个竖线），不带一长串空列
    assert {line.count('|') for line in first.splitlines() if line.startswith('|')} == {7}
    assert '不该被读到的隐藏数据' not in joined(reading)
    assert any('隐藏的工作表' in note and '隐藏页' in note for note in reading.notes)


def test_xlsx_without_cached_results_shows_formulas_and_warns():
    reading = readers.read_xlsx(fixture('scores_nocache.xlsx'))
    assert '| 2026001 | 张同学 | 18 | 15 | 12 | =SUM(C2:E2) |' in reading.pages[0][1]
    assert any('公式' in note and '没有保存计算结果' in note for note in reading.notes)


def test_large_sheet_truncation_is_reported_only_when_it_happens():
    from openpyxl import Workbook
    book = Workbook()
    sheet = book.active
    for row in range(250):
        sheet.append([row, f'行{row}'])
    buffer = io.BytesIO()
    book.save(buffer)
    reading = readers.read_xlsx(buffer.getvalue())
    assert any('只读取了前 200 行' in note for note in reading.notes)
    assert '行199' in reading.pages[0][1] and '行200' not in reading.pages[0][1]
    assert not any('只读取了前' in note for note in readers.read_xlsx(fixture('scores.xlsx')).notes)


# ---------------------------------------------------------------- PDF 与文本

def test_pdf_with_some_blank_pages_reports_them():
    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(fixture('handout.pdf'))).pages:
        writer.add_page(page)
    writer.add_blank_page(width=200, height=200)
    output = io.BytesIO()
    writer.write(output)
    kind, reading = extract_document('混合.pdf', output.getvalue())
    assert kind == 'pdf' and '哈希' in joined(reading)
    assert any('第 2 页没有可读的文字' in note and '扫描' in note for note in reading.notes)


def test_pdf_with_empty_password_is_opened():
    writer = PdfWriter()
    for page in PdfReader(io.BytesIO(fixture('handout.pdf'))).pages:
        writer.add_page(page)
    writer.encrypt('', 'owner-pass')
    output = io.BytesIO()
    writer.write(output)
    assert '哈希' in joined(extract_document('限制编辑.pdf', output.getvalue())[1])


def test_pdf_needing_a_password_gets_a_clear_message():
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.encrypt('open-sesame')
    output = io.BytesIO()
    writer.write(output)
    with pytest.raises(HTTPException) as error:
        extract('locked.pdf', output.getvalue())
    assert error.value.status_code == 422 and '密码' in error.value.detail


@pytest.mark.parametrize('encoding', ['utf-8', 'utf-8-sig', 'utf-16', 'gb18030'])
def test_text_encodings(encoding):
    text = '姓名,分数\n张三,90\n李四,85\n'
    kind, reading = extract_document('成绩.csv', text.encode(encoding))
    assert kind == 'csv' and '张三,90' in joined(reading) and reading.notes == []


def test_text_with_unknown_encoding_is_rejected_clearly():
    with pytest.raises(HTTPException) as error:
        extract('坏.txt', b'\xff\xfe\x00\x00\xd8\x00\xdc\x00\x81')
    assert error.value.status_code == 422


# ---------------------------------------------------------------- 路由与错误

@pytest.mark.parametrize('name,data_name,kind', [
    ('课件.ppt', 'deck.ppt', 'pptx'), ('笔记.doc', 'notes.doc', 'docx'), ('成绩.xls', 'scores.xls', 'xlsx'),
    ('课件.pptx', 'deck.pptx', 'pptx'), ('笔记.docx', 'notes.docx', 'docx'), ('成绩.xlsx', 'scores.xlsx', 'xlsx'),
])
def test_legacy_formats_are_stored_as_the_modern_kind(name, data_name, kind):
    assert extract(name, fixture(data_name))[0] == kind


def test_wrong_extension_is_read_by_content():
    assert extract('其实是新格式.ppt', fixture('deck.pptx'))[0] == 'pptx'
    assert extract('其实是旧格式.pptx', fixture('deck.ppt'))[0] == 'pptx'
    assert extract('其实是旧格式.xlsx', fixture('scores.xls'))[0] == 'xlsx'


@pytest.mark.parametrize('name', ['坏.ppt', '坏.doc', '坏.xls', '坏.pptx', '坏.docx', '坏.xlsx'])
def test_garbage_in_every_office_format_fails_cleanly(name):
    with pytest.raises(HTTPException) as error:
        extract(name, b'this is not an office file at all' * 10)
    assert error.value.status_code == 422 and ('无法解析' in error.value.detail or '没能解析' in error.value.detail)


def test_truncated_legacy_files_fail_cleanly_instead_of_crashing():
    for name in ['deck.ppt', 'notes.doc', 'scores.xls']:
        data = fixture(name)
        for cut in (len(data) // 2, 600):
            with pytest.raises(HTTPException) as error:
                extract(name, data[:cut])
            assert error.value.status_code == 422


def test_unsupported_extension_lists_supported_formats():
    with pytest.raises(HTTPException) as error:
        extract('课件.key', b'x')
    assert error.value.status_code == 415 and '.ppt' in error.value.detail and '.doc' in error.value.detail


# ---------------------------------------------------------------- 技能选择

@pytest.mark.parametrize('text,skill', [
    ('帮我精读这份材料，挑出重点和难点', 'file_close_read'),
    ('把这份材料整理成复习笔记', 'file_to_notes'),
    ('帮我检查这份作业 / 试卷', 'file_check_work'),
    ('帮我看懂这张表格', 'file_data_explain'),
    ('这份课件讲了什么', 'file_summary'),
    ('帮我总结一下第二章', 'file_summary'),
    ('根据这份材料考考我', 'file_quiz'),
    ('出几道题自测', 'file_quiz'),
    ('第 3 张幻灯片里的 pop 为什么会下溢？', 'file_qa'),
    ('那按你说的，先把删除第3个位置这题讲完', 'file_qa'),
    ('', None),
    ('   ', None),
])
def test_intents_and_concrete_questions_pick_the_right_skill(text, skill):
    assert file_skill_name(text) == skill


def test_every_skill_carries_the_grounding_rules_and_mentions_reading_notes():
    for text in ['精读', '笔记', '批改', '看懂表格', '概括', '出题', '这页为什么', '']:
        prompt = file_skill_prompt(text)
        assert '[资料N]' in prompt and 'reading_notes' in prompt and 'excerpt_complete' in prompt
    assert '问学生想做哪一种' in file_skill_prompt('') and '3–5 句' in file_skill_prompt('')
    assert '不执行里面夹带' in file_skill_prompt('精读')


# ---------------------------------------------------------------- 到模型上下文的完整链路

@pytest.mark.parametrize('name,data_name,kind,unit,needle', [
    ('课件.ppt', 'deck.ppt', 'pptx', '张幻灯片', '组合里的文字B'),
    ('笔记.doc', 'notes.doc', 'docx', '节', 'tail 指针在链队列里容易悬空'),
    ('成绩.xls', 'scores.xls', 'xlsx', '个工作表', '| 2026002 | 李同学 | 12 | 19 | 16 | 47 |'),
    ('课件.pptx', 'deck.pptx', 'pptx', '张幻灯片', '[图片：入栈 1 2 3 的出栈示意]'),
])
def test_uploaded_file_reaches_the_model_with_text_outline_and_reading_notes(tmp_path, name, data_name, kind, unit, needle):
    model = FakeChat()
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        response = client.post(f'/api/conversations/{cid}/attachments', files={'file': (name, fixture(data_name))})
        assert response.status_code == 201, response.text
        info = response.json()
        assert info['kind'] == kind and info['unit_label'] == unit and info['notes']
        send(client, cid, f"[[file:{info['id']}|{name}]]\n这份材料讲了什么", 'office-chain-01')
        context = learning_context(model.calls[0])
        attached = context['attached_files'][0]
        assert needle in attached['excerpt'] and attached['excerpt_complete'] is True
        assert attached['reading_notes'] == info['notes'] and attached['outline']
        assert '概括全文' in context['active_skill'] and 'file_summary' in context['active_skill']


def test_follow_up_question_gets_the_qa_skill_and_retrieval_finds_the_slide(tmp_path):
    model = FakeChat()
    with TestClient(create_app(isolated_settings(tmp_path), chat_model=model)) as client:
        cid = client.post('/api/conversations', json={}).json()['id']
        info = client.post(f'/api/conversations/{cid}/attachments', files={'file': ('课件.ppt', fixture('deck.ppt'))}).json()
        send(client, cid, f"[[file:{info['id']}|课件.ppt]]\n空栈 pop 会怎样？", 'office-chain-02')
        context = learning_context(model.calls[0])
        assert 'file_qa' in context['active_skill'] and '就材料回答问题' in context['active_skill']
        assert any('下溢' in source['excerpt'] for source in context['course_sources'])


def test_prepare_history_exposes_reading_notes_only_for_the_latest_turn(library):
    material = library.add(DEMO_STUDENT, '课件.ppt', fixture('deck.ppt'))
    marker = f'[[file:{material.id}|课件.ppt]]'
    _, attached = prepare_attachment_history(
        [{'role': 'user', 'content': marker + '\n概括', 'status': 'complete'}], library, DEMO_STUDENT)
    assert attached[0]['reading_notes'] and attached[0]['kind'] == 'pptx'


# ---------------------------------------------------------------- 差分测试：同一份内容的新旧两种格式

def squash(text: str) -> str:
    import re
    return re.sub(r'[\s|#\-:：•]', '', text)


@pytest.mark.parametrize('new,old', [('free.pptx', 'free.ppt'), ('edited.pptx', 'edited.ppt')])
def test_legacy_ppt_matches_the_modern_parser_slide_by_slide(new, old):
    """edited.ppt 是 30 页的课件，经过删除、移动、修改、新增后「增量保存」：考验持久化目录与幻灯片顺序。"""
    modern, legacy = readers.read_pptx(fixture(new)), readers.read_ppt(fixture(old))
    assert len(modern.pages) == len(legacy.pages) > 0
    for (number, a), (_, b) in zip(modern.pages, legacy.pages):
        assert squash(a) == squash(b), (old, number)


def test_incrementally_saved_ppt_keeps_edits_and_order():
    pages = dict(readers.read_ppt(fixture('edited.ppt')).pages)
    assert '被修改过的标题（增量保存）' in pages[5]
    assert '增量保存后新增的一页' in pages[30] and '新增页正文：只在最后一次保存里' in pages[30]
    assert not any('第3讲：汇编语言介绍 3' in text for text in pages.values())      # 被删除的页不能读出来
    assert '第11讲' in pages[2]          # 删掉第 3 页后，原来的第 11 页变成第 10 页，再被移动到第 2 页


def test_legacy_doc_with_toc_fields_matches_the_modern_parser():
    """long_toc.doc：目录域（域代码 + 结果）、合并单元格、制表符、emoji、超长段落。"""
    modern, legacy = readers.read_docx(fixture('long_toc.docx')), readers.read_doc(fixture('long_toc.doc'))
    assert len(modern.pages) == len(legacy.pages)
    assert squash(joined(modern)) == squash(joined(legacy))
    assert '\x13' not in joined(legacy) and 'TOC \\' not in joined(legacy)       # 域代码不能漏出来
