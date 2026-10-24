"""课件里的图片：取图、过滤装饰图、识图并回对应幻灯片。识图模型一律用模拟，不访问网络。

tests/fixtures/office/pictures.ppt / .pptx 由真实 PowerPoint 生成：第 1-4 页都放了同一个小 logo，
第 2 页是文字 + 一张 PNG 截图，第 3 页只有一张 JPEG 图，第 4 页只有文字。
"""

import asyncio
import io
import json
from pathlib import Path

import httpx
import pytest
from fastapi import HTTPException
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches
from pydantic import SecretStr

from types import SimpleNamespace

from app.agent_tools import documents
from app.memory.store import Store
from app.services.file_store import FileStore
from app.services.materials import MaterialStore
from app.services.office_readers import Reading
from app.services.slide_images import MAX_PICTURES, SLIDE_PROMPT, add_picture_text, collect_pictures
from app.vision.settings import VisionConfig

FIXTURES = Path(__file__).parent / "fixtures" / "office"


def fixture(name: str) -> bytes:
    return (FIXTURES / name).read_bytes()


class Models:
    def __init__(self, configured=True):
        self.configured = configured

    def config(self):
        return None

    def vision_config(self, _):
        if not self.configured:
            return VisionConfig()
        return VisionConfig(base_url="https://vision.test/v1", model="vlm", api_key=SecretStr("k"), source="env")


def reply(text):
    return httpx.Response(200, json={"choices": [{"message": {"content": text}}]})


def run(coroutine):
    return asyncio.run(coroutine)


def text_only(kind="pptx"):
    return Reading([(1, "# 标题"), (2, "# 变量存储分配\n下图是数据段里变量的存放顺序"), (3, ""), (4, "# 小结")],
                   ["2 张幻灯片含图片（共 6 张），图片里的文字没有读取。"])


@pytest.mark.parametrize("name", ["pictures.pptx", "pictures.ppt"])
def test_pictures_are_taken_per_slide_and_the_repeated_logo_is_ignored(name):
    found = collect_pictures("pptx", fixture(name))
    assert found.total == 6
    assert [picture.slide for picture in found.pictures] == [2, 3]
    assert found.decorative == 4 and found.unsupported == 0 and found.over_limit == 0
    assert all(picture.data.startswith(b"\x89PNG") for picture in found.pictures)


@pytest.mark.parametrize("name", ["pictures.pptx", "pictures.ppt"])
def test_picture_text_is_merged_into_the_right_slides_and_notes_replace_the_old_one(name):
    seen = []

    def handler(request):
        body = json.loads(request.content)
        seen.append(body["messages"][0]["content"][1]["text"])
        index = len(seen)
        return reply(f"图中文字 {index}")

    reading = run(add_picture_text("pptx", fixture(name), text_only(), Models(), transport=httpx.MockTransport(handler)))
    assert seen == [SLIDE_PROMPT, SLIDE_PROMPT]
    pages = dict(reading.pages)
    assert "[图片内容]" in pages[2] and pages[2].startswith("# 变量存储分配")
    assert "[图片内容]" in pages[3] and pages[1] == "# 标题" and pages[4] == "# 小结"
    assert any("第 2、3 页的图片由识图模型识别" in note for note in reading.notes)
    assert not any("图片里的文字没有读取" in note for note in reading.notes)


def test_without_a_vision_model_the_text_is_kept_and_the_note_says_what_to_do():
    reading = run(add_picture_text("pptx", fixture("pictures.pptx"), text_only(), Models(configured=False)))
    assert dict(reading.pages) == dict(text_only().pages)
    assert any("2 张内容图片" in note and "配置" in note for note in reading.notes)
    assert not any("共 6 张" in note for note in reading.notes)


def test_a_failed_picture_is_reported_and_does_not_lose_the_others_or_the_text():
    calls = []

    def handler(request):
        calls.append(1)
        if len(calls) == 1:
            return httpx.Response(500, json={"error": "boom"})
        return reply("第二张的文字")

    reading = run(add_picture_text("pptx", fixture("pictures.pptx"), text_only(), Models(), transport=httpx.MockTransport(handler)))
    pages = dict(reading.pages)
    assert sum("[图片内容]" in text for text in pages.values()) == 1
    assert pages[1] == "# 标题"
    assert any("识别失败" in note for note in reading.notes)


def test_decks_without_pictures_are_returned_unchanged():
    # deck.pptx only has one tiny 120x80 swatch: nothing worth sending to a vision model
    deck = fixture("deck.pptx")
    assert collect_pictures("pptx", deck).pictures == []
    original = Reading([(1, "x")], ["原有说明"])
    assert run(add_picture_text("pptx", deck, original, Models())) is original


def test_picture_count_is_capped_and_small_icons_are_dropped():
    presentation = Presentation()
    for number in range(MAX_PICTURES + 4):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        image = Image.new("RGB", (400, 300), (255, 255, 255))
        ImageDraw.Draw(image).text((20, 20), f"picture {number}", fill=(0, 0, 0))
        buffer = io.BytesIO()
        image.save(buffer, "PNG")
        buffer.seek(0)
        slide.shapes.add_picture(buffer, Inches(1), Inches(1))
        icon = io.BytesIO()
        Image.new("RGB", (32, 32), (10, 10, 10)).save(icon, "PNG")
        icon.seek(0)
        slide.shapes.add_picture(icon, Inches(0.2), Inches(0.2))
    output = io.BytesIO()
    presentation.save(output)
    found = collect_pictures("pptx", output.getvalue())
    assert len(found.pictures) == MAX_PICTURES and found.over_limit == 4
    assert found.decorative >= MAX_PICTURES + 4  # every 32px icon (also repeated on every slide)


@pytest.fixture
def library(tmp_path):
    store = Store(tmp_path / "slides.sqlite3")
    store.initialize()
    result = MaterialStore(store)
    result.initialize()
    return result


def image_only_deck() -> bytes:
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    image = Image.new("RGB", (600, 400), "white")
    ImageDraw.Draw(image).text((30, 30), "scanned slide", fill="black")
    buffer = io.BytesIO()
    image.save(buffer, "PNG")
    buffer.seek(0)
    picture = slide.shapes.add_picture(buffer, Inches(1), Inches(1))
    picture._element.nvPicPr.cNvPr.set("descr", "")  # python-pptx fills it with the file name, which reads as text
    output = io.BytesIO()
    presentation.save(output)
    return output.getvalue()


def test_upload_uses_the_slide_hook_and_accepts_decks_that_are_only_pictures(library):
    async def hook(kind, data, reading):
        return await add_picture_text(kind, data, reading, Models(), transport=httpx.MockTransport(lambda _: reply("讲义第一页的文字")))

    # without the hook a deck of pure pictures has no text and is refused, as before
    with pytest.raises(HTTPException) as refused:
        run(library.add_with_ocr("student", "扫描.pptx", image_only_deck()))
    assert refused.value.status_code == 422

    library.slide_ocr = hook
    material = run(library.add_with_ocr("student", "扫描.pptx", image_only_deck()))
    assert material.kind == "pptx" and material.chunk_count == 1
    attachment = library.attachment(material.id, "student", excerpt=True)
    assert "讲义第一页的文字" in attachment["excerpt"]
    assert any("由识图模型识别" in note for note in attachment["notes"])

    # legacy .ppt goes through the same hook, and a broken hook never loses the text
    async def broken(kind, data, reading):
        raise RuntimeError("vision is down")

    library.slide_ocr = broken
    legacy = run(library.add_with_ocr("student", "课件.ppt", fixture("pictures.ppt")))
    assert legacy.chunk_count > 0
    assert "变量存储分配" in library.attachment(legacy.id, "student", excerpt=True)["excerpt"]


@pytest.fixture
def tool_env(tmp_path):
    store = Store(tmp_path / "tool.sqlite3")
    store.initialize()
    files = FileStore(store)
    files.initialize()
    materials = MaterialStore(store)
    materials.initialize()
    materials.originals = files
    return materials, SimpleNamespace(files=files, library=materials, student_id="student", generated=[])


def test_read_document_returns_the_picture_text_that_was_recognized_at_upload(tool_env):
    materials, ctx = tool_env

    async def hook(kind, data, reading):
        return await add_picture_text(kind, data, reading, Models(), transport=httpx.MockTransport(lambda _: reply("DS:0002  0915H")))

    materials.slide_ocr = hook
    material = run(materials.add_with_ocr("student", "课件.ppt", fixture("pictures.ppt")))
    output = documents.read_document({"file_id": material.id, "start": 2, "limit": 2}, ctx)
    assert "【第 2 页】" in output and "变量存储分配" in output
    assert "[图片内容]\nDS:0002  0915H" in output           # slide 2 text + its picture text
    assert output.count("DS:0002  0915H") == 2                 # slide 3 is only a picture
    assert "由识图模型识别" in output and "没读到的内容" in output


def test_read_document_without_recognized_text_is_unchanged(tool_env):
    materials, ctx = tool_env
    material = materials.add("student", "课件.pptx", fixture("pictures.pptx"))
    output = documents.read_document({"file_id": material.id, "start": 2, "limit": 1}, ctx)
    assert "变量存储分配" in output and "[图片内容]" not in output


def test_read_document_reads_a_scanned_pdf_from_its_stored_ocr_text(tool_env):
    materials, ctx = tool_env
    image = Image.new("RGB", (800, 1100), "white")
    ImageDraw.Draw(image).text((60, 60), "scanned", fill="black")
    out = io.BytesIO()
    image.save(out, "PDF", resolution=100)

    async def ocr(_data):
        return Reading([(1, "扫描讲义第一页")], ["这份 PDF 是扫描图片，文字由识图模型识别。"], extra={1: "扫描讲义第一页"})

    materials.ocr = ocr
    material = run(materials.add_with_ocr("student", "讲义.pdf", out.getvalue()))
    output = documents.read_document({"file_id": material.id}, ctx)
    assert "扫描讲义第一页" in output and "由识图模型识别" in output

