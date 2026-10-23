"""Office 文件读取：.pptx / .docx / .xlsx 以及旧二进制格式 .ppt / .doc / .xls。

每个读取函数返回 Reading：按页 / 节 / 幻灯片 / 工作表切好的文字，加上「没读到什么」的说明。
说明会交给对话模型，让它如实告诉学生哪些内容（图片里的字、隐藏页等）没有读到，而不是当成没有。
"""
import io
import re
import struct
from dataclasses import dataclass, field
from datetime import date, datetime

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import HTTPException
from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

MAX_PAGES = 200
MAX_SHEET_ROWS = 200
MAX_SHEET_COLUMNS = 30
MAX_CELL_CHARS = 200
DOCX_CHUNK = 1500

_NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'dgm': 'http://schemas.openxmlformats.org/drawingml/2006/diagram',
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
}
_CONTROL = re.compile(r'[\x00-\x08\x0e-\x1f\x7f]')


@dataclass
class Reading:
    pages: list[tuple[int, str]]
    notes: list[str] = field(default_factory=list)


def _x(element, expression: str):
    """带命名空间的 XPath（python-pptx / python-docx 自带的 xpath 不接受 namespaces 参数）。"""
    return etree.XPath(expression, namespaces=_NS)(element)


def _tidy(text: str) -> str:
    text = _CONTROL.sub('', text.replace('\r\n', '\n').replace('\r', '\n').replace('\x0b', '\n').replace('\x0c', '\n'))
    text = re.sub(r'[ \t]+\n', '\n', text)
    return re.sub(r'\n{3,}', '\n\n', text).strip()


def _row(values) -> str:
    return '| ' + ' | '.join(str(v if v is not None else '').replace('|', '\\|').replace('\n', '<br>')
                             for v in values) + ' |'


def _table_lines(rows: list[list[str]]) -> list[str]:
    rows = [row for row in rows if any(str(cell).strip() for cell in row)]
    if not rows:
        return []
    width = max(len(row) for row in rows)
    rows = [[*row, *([''] * (width - len(row)))] for row in rows]
    return [_row(rows[0]), _row(['---'] * width), *(_row(row) for row in rows[1:])]


def _numbers(values: list[int]) -> str:
    values = sorted(set(values))
    shown = '、'.join(str(v) for v in values[:12])
    return f'第 {shown}{" 等" if len(values) > 12 else ""}'


def _limit_pages(pages, unit: str):
    if len(pages) > MAX_PAGES:
        raise HTTPException(413, f'单份资料最多 {MAX_PAGES} {unit}，请按章节拆分。')


# ------------------------------------------------------------------ pptx

def _smartart_text(shape, part) -> list[str]:
    """SmartArt（流程图、层次结构图）的文字不在形状里，而在单独的 diagram data 部件里。"""
    try:
        rid = _x(shape._element, './/dgm:relIds/@r:dm')
        if not rid:
            return []
        blob = part.related_part(rid[0]).blob
        root = etree.fromstring(blob)
        points = _x(root, './/dgm:pt[not(@type) or @type="node"]')
        lines = []
        for point in points:
            text = ''.join(_x(point, './/a:t/text()')).strip()
            if text:
                lines.append(text)
        return lines
    except Exception:  # noqa: BLE001 - SmartArt 读不出来不能让整份课件失败
        return []


def _chart_line(shape) -> str:
    try:
        chart = shape.chart
        parts = []
        if chart.has_title and chart.chart_title.has_text_frame and chart.chart_title.text_frame.text.strip():
            parts.append(chart.chart_title.text_frame.text.strip())
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories][:12]
        series = [s.name for s in plot.series if s.name][:6]
        if series:
            parts.append('系列：' + '、'.join(series))
        if categories:
            parts.append('类别：' + '、'.join(categories))
        return '[图表' + ('：' + '；'.join(parts) if parts else '') + ']'
    except Exception:  # noqa: BLE001
        return '[图表]'


def _shape_lines(shapes, part, skip_id, stats: dict) -> list[str]:
    lines: list[str] = []

    def position(shape):
        try:
            return (shape.top or 0, shape.left or 0)
        except Exception:  # noqa: BLE001
            return (0, 0)

    for shape in sorted(shapes, key=position):
        if skip_id is not None and shape.shape_id == skip_id:
            continue
        kind = getattr(shape, 'shape_type', None)
        if kind == MSO_SHAPE_TYPE.GROUP:
            lines.extend(_shape_lines(shape.shapes, part, None, stats))
            continue
        if kind == MSO_SHAPE_TYPE.PICTURE:
            stats['images'] += 1
            alt = ''.join(_x(shape._element, './/p:cNvPr/@descr')).strip()
            if alt and not re.fullmatch(r'(?i)(picture|image|图片)\s*\d*', alt):
                lines.append(f'[图片：{alt[:100]}]')
            continue
        if getattr(shape, 'has_chart', False) and shape.has_chart:
            stats['charts'] += 1
            lines.append(_chart_line(shape))
            continue
        if kind in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
            stats['objects'] += 1
            continue
        smart = _smartart_text(shape, part) if getattr(shape, 'has_text_frame', True) is False or kind is None else []
        if smart:
            lines.extend('- ' + line for line in smart)
            continue
        if getattr(shape, 'has_table', False) and shape.has_table:
            lines.extend(_table_lines([[cell.text for cell in row.cells] for row in shape.table.rows]))
            continue
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.replace('\x0b', ' ').strip()
                if not text:
                    continue
                level = max(0, min(int(paragraph.level or 0), 4))
                lines.append(('  ' * (level - 1) + '- ' + text) if level else text)
    return lines


def read_pptx(data: bytes) -> Reading:
    presentation = Presentation(io.BytesIO(data))
    slides = list(presentation.slides)
    _limit_pages(slides, '张幻灯片')
    pages, notes = [], []
    stats_all = {'images': 0, 'charts': 0, 'objects': 0}
    image_slides = []
    for number, slide in enumerate(slides, 1):
        stats = {'images': 0, 'charts': 0, 'objects': 0}
        title = slide.shapes.title
        parts = ['# ' + title.text.strip()] if title is not None and title.text.strip() else []
        parts.extend(_shape_lines(slide.shapes, slide.part, title.shape_id if title is not None else None, stats))
        if slide._element.get('show') == '0':
            parts.append('（这张幻灯片在放映时是隐藏的）')
        if slide.has_notes_slide:
            text = (slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame is not None else '').strip()
            if text:
                parts.append('备注：' + text)
        if stats['images']:
            image_slides.append(number)
        for key in stats_all:
            stats_all[key] += stats[key]
        pages.append((number, '\n'.join(parts)))
    if image_slides:
        notes.append(f'{_numbers(image_slides)} 张幻灯片含图片（共 {stats_all["images"]} 张），图片里的文字没有读取。')
    if stats_all['objects']:
        notes.append(f'含 {stats_all["objects"]} 个嵌入的 Office 对象，内容没有读取。')
    return Reading(pages, notes)


# ------------------------------------------------------------------ docx

def _docx_part_text(package, partname: str, tag: str) -> list[str]:
    for part in package.iter_parts():
        if str(part.partname) != partname:
            continue
        root = etree.fromstring(part.blob)
        lines = []
        for node in _x(root, f'.//w:{tag}[not(@w:type)]'):
            text = ' '.join(t.strip() for paragraph in _x(node, './/w:p')
                            for t in [''.join(_x(paragraph, './/w:t/text()'))] if t.strip())
            if text.strip():
                lines.append(text.strip())
        return lines
    return []


def _split_blocks(blocks: list[str], size: int = DOCX_CHUNK) -> list[tuple[int, str]]:
    pages, buffer = [], ''
    for block in blocks:
        for start in range(0, len(block), size):
            part = block[start:start + size]
            if buffer and len(buffer) + len(part) + 1 > size:
                pages.append((len(pages) + 1, buffer))
                buffer = ''
            buffer += ('\n' if buffer else '') + part
    if buffer:
        pages.append((len(pages) + 1, buffer))
    return pages


def read_docx(data: bytes) -> Reading:
    document = Document(io.BytesIO(data))
    blocks: list[str] = []
    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            text = block.text.strip()
            # 文本框里的字不在段落正文里
            extra = [t.strip() for t in _x(block._p, './/w:txbxContent[not(ancestor::mc:Fallback)]//w:t/text()') if t.strip()]
            style = block.style.name if block.style else ''
            heading = re.match(r'(?:Heading|标题)\s*([1-9])$', style, re.I)
            if text and (heading or style == 'Title'):
                text = '#' * (min(int(heading[1]), 6) if heading else 1) + ' ' + text
            elif text and (style.startswith('List') or style.startswith('列表') or _x(block._p, './w:pPr/w:numPr')):
                text = '- ' + text
            if extra:
                text = (text + '\n' if text else '') + '[文本框] ' + ' '.join(extra)
        elif isinstance(block, Table):
            rows = []
            for row in block.rows:
                cells, previous = [], None
                for cell in row.cells:
                    if cell._tc is previous:  # 横向合并的单元格会重复出现
                        continue
                    previous = cell._tc
                    cells.append(cell.text.strip())
                rows.append(cells)
            text = '\n'.join(_table_lines(rows))
        else:
            continue
        if text:
            blocks.append(text)
    notes: list[str] = []
    headers = []
    for section in document.sections:
        for part in (section.header, section.footer, section.first_page_header, section.first_page_footer,
                     section.even_page_header, section.even_page_footer):
            try:
                text = ' '.join(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:  # noqa: BLE001
                continue
            if text and text not in headers:
                headers.append(text)
    package = document.part.package
    footnotes = _docx_part_text(package, '/word/footnotes.xml', 'footnote')
    endnotes = _docx_part_text(package, '/word/endnotes.xml', 'endnote')
    if footnotes:
        blocks.append('【脚注】\n' + '\n'.join(f'- {line}' for line in footnotes))
    if endnotes:
        blocks.append('【尾注】\n' + '\n'.join(f'- {line}' for line in endnotes))
    if headers:
        blocks.append('【页眉页脚】' + ' / '.join(headers))
    images = len(document.inline_shapes)
    if images:
        notes.append(f'文中含 {images} 张图片，图片里的文字没有读取。')
    pages = _split_blocks(blocks)
    _limit_pages(pages, '节')
    return Reading(pages, notes)


# ------------------------------------------------------------------ xlsx / xls

def _plain(value) -> str:
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
        return str(int(value))
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d %H:%M' if (value.hour or value.minute) else '%Y-%m-%d')
    if isinstance(value, date):
        return value.strftime('%Y-%m-%d')
    return str(value).strip()[:MAX_CELL_CHARS]


def _trim_columns(rows: list[list[str]]) -> list[list[str]]:
    """按实际用到的列数裁掉右边整列为空的部分（只读模式会把每行补到 max_col）。"""
    width = max((index + 1 for row in rows for index, cell in enumerate(row) if cell), default=0)
    return [row[:width] for row in rows]


def _sheet_page(title: str, total_rows, total_cols, rows: list[list[str]], clipped_cols: bool) -> str:
    rows = _trim_columns(rows)
    parts = [f'# {title}', f'总行数：{total_rows}；总列数：{total_cols}；仅展示前 {MAX_SHEET_ROWS} 行'
             + (f'、前 {MAX_SHEET_COLUMNS} 列' if clipped_cols else '') + '。']
    parts.extend(_table_lines(rows) if any(cell for row in rows for cell in row) else ['预览范围内无数据。'])
    return '\n'.join(parts)


def read_xlsx(data: bytes) -> Reading:
    values_book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    formulas_book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    pages, notes, hidden, uncached, clipped = [], [], [], 0, False
    try:
        sheets = list(values_book.worksheets)
        _limit_pages(sheets, '个工作表')
        has_values = False
        for number, sheet in enumerate(sheets, 1):
            if sheet.sheet_state != 'visible':
                hidden.append(sheet.title)
                pages.append((number, f'# {sheet.title}\n（隐藏的工作表，没有读取）'))
                continue
            if sheet.max_row is None or sheet.max_column is None:
                sheet.calculate_dimension(force=True)
            formula_rows = list(formulas_book[sheet.title].iter_rows(
                max_row=min(sheet.max_row or 0, MAX_SHEET_ROWS), max_col=MAX_SHEET_COLUMNS, values_only=True))
            rows = []
            for r_index, row in enumerate(sheet.iter_rows(
                    max_row=min(sheet.max_row or 0, MAX_SHEET_ROWS), max_col=MAX_SHEET_COLUMNS, values_only=True)):
                cells = []
                for c_index, value in enumerate(row):
                    if value is None:
                        formula = formula_rows[r_index][c_index] if r_index < len(formula_rows) and c_index < len(formula_rows[r_index]) else None
                        if isinstance(formula, str) and formula.startswith('='):
                            uncached += 1
                            value = formula  # 公式没有缓存结果时，显示公式本身，不当成 0
                    cells.append(_plain(value))
                rows.append(cells)
            if any(cell for row in rows for cell in row):
                has_values = True
            clipped = clipped or (sheet.max_row or 0) > MAX_SHEET_ROWS or (sheet.max_column or 0) > MAX_SHEET_COLUMNS
            pages.append((number, _sheet_page(sheet.title, sheet.max_row, sheet.max_column, rows,
                                              (sheet.max_column or 0) > MAX_SHEET_COLUMNS)))
        if not has_values:
            raise HTTPException(422, '没有提取到文字，工作表前 200 行内没有可读取的数据。')
    finally:
        values_book.close()
        formulas_book.close()
    if hidden:
        notes.append(f'{len(hidden)} 个隐藏的工作表没有读取：{"、".join(hidden[:6])}。')
    if uncached:
        notes.append(f'有 {uncached} 个公式单元格没有保存计算结果（文件不是用 Excel 保存的），这里显示的是公式本身，不是数值。')
    if clipped:
        notes.append(f'有的工作表超过 {MAX_SHEET_ROWS} 行或 {MAX_SHEET_COLUMNS} 列，只读取了前 {MAX_SHEET_ROWS} 行、前 {MAX_SHEET_COLUMNS} 列。')
    return Reading(pages, notes)


def read_xls(data: bytes) -> Reading:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data)
    except xlrd.XLRDError as exc:
        text = str(exc)
        if 'encrypt' in text.lower() or 'password' in text.lower():
            raise HTTPException(422, '这份 Excel 文件有密码保护，请先解除密码。') from None
        raise
    pages, hidden, has_values, clipped = [], [], False, False
    sheets = [book.sheet_by_index(i) for i in range(book.nsheets)]
    _limit_pages(sheets, '个工作表')
    for number, sheet in enumerate(sheets, 1):
        if sheet.visibility != 0:
            hidden.append(sheet.name)
            pages.append((number, f'# {sheet.name}\n（隐藏的工作表，没有读取）'))
            continue
        rows = []
        for r in range(min(sheet.nrows, MAX_SHEET_ROWS)):
            cells = []
            for c in range(min(sheet.ncols, MAX_SHEET_COLUMNS)):
                cell = sheet.cell(r, c)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    try:
                        cells.append(_plain(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode)))
                        continue
                    except Exception:  # noqa: BLE001
                        pass
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    cells.append('')
                elif cell.ctype == xlrd.XL_CELL_ERROR:
                    cells.append('#错误')
                else:
                    cells.append(_plain(cell.value))
            rows.append(cells)
        if any(cell for row in rows for cell in row):
            has_values = True
        clipped = clipped or sheet.nrows > MAX_SHEET_ROWS or sheet.ncols > MAX_SHEET_COLUMNS
        pages.append((number, _sheet_page(sheet.name, sheet.nrows, sheet.ncols, rows, sheet.ncols > MAX_SHEET_COLUMNS)))
    if not has_values:
        raise HTTPException(422, '没有提取到文字，工作表前 200 行内没有可读取的数据。')
    notes = []
    if hidden:
        notes.append(f'{len(hidden)} 个隐藏的工作表没有读取：{"、".join(hidden[:6])}。')
    if clipped:
        notes.append(f'有的工作表超过 {MAX_SHEET_ROWS} 行或 {MAX_SHEET_COLUMNS} 列，只读取了前 {MAX_SHEET_ROWS} 行、前 {MAX_SHEET_COLUMNS} 列。')
    return Reading(pages, notes)


# ------------------------------------------------------------------ 旧版 .ppt

_RT_DOCUMENT, _RT_SLIDE, _RT_SLIDE_ATOM, _RT_NOTES = 1000, 1006, 1007, 1008
_RT_SLIDE_LIST, _RT_SLIDE_PERSIST, _RT_SSSLIDE_INFO = 4080, 1011, 1017
_RT_TEXT_HEADER, _RT_TEXT_CHARS, _RT_TEXT_BYTES = 3999, 4000, 4008
_RT_USER_EDIT, _RT_PERSIST_DIR = 4085, 6002
_TITLE_TYPES = {0, 6}   # 标题 / 居中标题
_NOTES_TYPE = 2


def _ole(data: bytes):
    import olefile
    if not olefile.isOleFile(data):
        raise ValueError('不是有效的 Office 二进制文件')
    return olefile.OleFileIO(io.BytesIO(data))


def _records(buf: bytes, start: int, end: int):
    """顺序读取 [start, end) 里的一层记录，产出 (类型, 版本, 实例, 数据起点, 长度)。"""
    offset = start
    while offset + 8 <= end:
        ver_inst, rtype, length = struct.unpack_from('<HHI', buf, offset)
        body = offset + 8
        if body + length > end:
            break
        yield rtype, ver_inst & 0xF, ver_inst >> 4, body, length
        offset = body + length


def _ppt_text(rtype: int, raw: bytes) -> str:
    text = raw.decode('utf-16-le', 'replace') if rtype == _RT_TEXT_CHARS else raw.decode('cp1252', 'replace')
    return _tidy(text)


def _ppt_collect(buf: bytes, start: int, end: int, out: list[tuple[int, str]]):
    """递归收集一个容器里所有文本块：[(文本类型, 文字), ...]，保持文档里的先后顺序。"""
    kind = None
    for rtype, ver, _inst, body, length in _records(buf, start, end):
        if ver == 0xF:
            _ppt_collect(buf, body, body + length, out)
        elif rtype == _RT_TEXT_HEADER and length >= 4:
            kind = struct.unpack_from('<I', buf, body)[0]
        elif rtype in (_RT_TEXT_CHARS, _RT_TEXT_BYTES):
            text = _ppt_text(rtype, buf[body:body + length])
            if text and text != '*':
                out.append((kind if kind is not None else 4, text))
            kind = None


def _persist_directory(ole, buf: bytes) -> dict[int, int]:
    """persistId → 流内偏移。沿着 UserEditAtom 链从新到旧合并，新的覆盖旧的。"""
    directory: dict[int, int] = {}
    try:
        current = ole.openstream('Current User').read()
        edit = struct.unpack_from('<I', current, 16)[0]
    except Exception:  # noqa: BLE001
        return directory
    seen = set()
    while edit and edit not in seen and edit + 8 <= len(buf):
        seen.add(edit)
        _ver, rtype, _length = struct.unpack_from('<HHI', buf, edit)
        if rtype != _RT_USER_EDIT:
            break
        body = edit + 8
        last_edit, persist_dir = struct.unpack_from('<II', buf, body + 8)
        if persist_dir + 8 <= len(buf):
            _, ptype, plen = struct.unpack_from('<HHI', buf, persist_dir)
            if ptype == _RT_PERSIST_DIR:
                pos, stop = persist_dir + 8, persist_dir + 8 + plen
                while pos + 4 <= stop:
                    head = struct.unpack_from('<I', buf, pos)[0]
                    pos += 4
                    first, count = head & 0xFFFFF, head >> 20
                    for index in range(count):
                        if pos + 4 > stop:
                            break
                        directory.setdefault(first + index, struct.unpack_from('<I', buf, pos)[0])
                        pos += 4
        edit = last_edit
    return directory


def read_ppt(data: bytes) -> Reading:
    ole = _ole(data)
    try:
        if not ole.exists('PowerPoint Document'):
            raise ValueError('没有找到幻灯片内容')
        has_pictures = ole.exists('Pictures') and ole.get_size('Pictures') > 0
        buf = ole.openstream('PowerPoint Document').read()
        directory = _persist_directory(ole, buf)
        # 幻灯片顺序 = Document 里 SlideListWithText(实例 0) 里 SlidePersistAtom 的顺序
        slide_ids: list[int] = []
        notes_persist: dict[int, int] = {}       # 备注页 slideId → persistId
        for rtype, _v, _i, body, length in _records(buf, 0, len(buf)):
            if rtype != _RT_DOCUMENT:
                continue
            for rtype2, _v2, inst2, body2, length2 in _records(buf, body, body + length):
                if rtype2 != _RT_SLIDE_LIST or inst2 not in (0, 2):
                    continue
                for rtype3, _v3, _i3, body3, length3 in _records(buf, body2, body2 + length2):
                    if rtype3 != _RT_SLIDE_PERSIST or length3 < 16:
                        continue
                    persist_id, slide_id = struct.unpack_from('<I', buf, body3)[0], struct.unpack_from('<I', buf, body3 + 12)[0]
                    if inst2 == 0:
                        slide_ids.append(persist_id)
                    else:
                        notes_persist[slide_id] = persist_id
        containers = [directory[i] for i in slide_ids if i in directory]
        if not containers:  # 没有持久化目录（或损坏）：按流里 Slide 容器出现的顺序读
            containers = [off - 8 for rtype, _v, _i, off, _l in _records(buf, 0, len(buf)) if rtype == _RT_SLIDE]
        _limit_pages(containers, '张幻灯片')
        pages = []
        for number, offset in enumerate(containers, 1):
            _ver, rtype, length = struct.unpack_from('<HHI', buf, offset)
            if rtype != _RT_SLIDE:
                pages.append((number, ''))
                continue
            texts: list[tuple[int, str]] = []
            _ppt_collect(buf, offset + 8, offset + 8 + length, texts)
            notes_text, is_hidden, notes_id = '', False, 0
            for rtype2, _v2, _i2, body2, length2 in _records(buf, offset + 8, offset + 8 + length):
                if rtype2 == _RT_SLIDE_ATOM and length2 >= 20:
                    notes_id = struct.unpack_from('<I', buf, body2 + 16)[0]
                elif rtype2 == _RT_SSSLIDE_INFO and length2 >= 12:
                    is_hidden = bool(struct.unpack_from('<H', buf, body2 + 10)[0] & 0x0004)
            if notes_id and notes_persist.get(notes_id) in directory:
                noff = directory[notes_persist[notes_id]]
                if noff + 8 <= len(buf) and struct.unpack_from('<H', buf, noff + 2)[0] == _RT_NOTES:
                    note_texts: list[tuple[int, str]] = []
                    _ppt_collect(buf, noff + 8, noff + 8 + struct.unpack_from('<I', buf, noff + 4)[0], note_texts)
                    notes_text = '\n'.join(t for k, t in note_texts if k == _NOTES_TYPE)
            lines = []
            for kind, text in texts:
                if kind == _NOTES_TYPE:
                    continue
                if kind in _TITLE_TYPES and not any(l.startswith('# ') for l in lines):
                    lines.append('# ' + text.replace('\n', ' '))
                else:
                    lines.extend(text.split('\n'))
            if is_hidden:
                lines.append('（这张幻灯片在放映时是隐藏的）')
            if notes_text:
                lines.append('备注：' + notes_text)
            pages.append((number, '\n'.join(line for line in lines if line.strip())))
    finally:
        ole.close()
    notes = ['这份课件含图片，图片、图表里的文字没有读取。'] if has_pictures else []
    return Reading(pages, notes)


# ------------------------------------------------------------------ 旧版 .doc

def _doc_pieces(word: bytes, table: bytes, fc_clx: int, lcb_clx: int) -> list[tuple[int, int, int, bool]]:
    clx = table[fc_clx:fc_clx + lcb_clx]
    pos = 0
    while pos < len(clx) and clx[pos] == 0x01:           # 跳过 Prc
        pos += 3 + struct.unpack_from('<H', clx, pos + 1)[0]
    if pos >= len(clx) or clx[pos] != 0x02:
        raise ValueError('找不到文字块表')
    length = struct.unpack_from('<I', clx, pos + 1)[0]
    plc = clx[pos + 5:pos + 5 + length]
    count = (length - 4) // 12
    cps = struct.unpack_from(f'<{count + 1}I', plc, 0)
    pieces = []
    for index in range(count):
        fc = struct.unpack_from('<I', plc, 4 * (count + 1) + 8 * index + 2)[0]
        compressed = bool(fc & 0x40000000)
        pieces.append((cps[index], cps[index + 1], (fc & 0x3FFFFFFF) // 2 if compressed else fc, compressed))
    return pieces


def _doc_text(word: bytes, pieces) -> str:
    chunks = []
    for cp_start, cp_end, offset, compressed in pieces:
        chars = cp_end - cp_start
        if compressed:
            chunks.append(word[offset:offset + chars].decode('cp1252', 'replace'))
        else:
            chunks.append(word[offset:offset + chars * 2].decode('utf-16-le', 'replace'))
    return ''.join(chunks)


def _doc_clean(raw: str) -> tuple[str, int]:
    """去掉域代码（只留结果）、把单元格标记整理成表格行，返回 (文字, 图片占位数)。"""
    images = raw.count('\x01')
    out, depth, show = [], 0, []
    for ch in raw:
        if ch == '\x13':
            depth += 1
            show.append(False)
        elif ch == '\x14' and show:
            show[-1] = True
        elif ch == '\x15' and show:
            show.pop()
            depth -= 1
        elif depth == 0 or all(show):
            out.append(ch)
    text = ''.join(out)
    text = text.replace('\x07\x07', '\n\x00ROW\x00').replace('\r\x07', ' | ').replace('\x07', ' | ')
    text = text.replace('\x00ROW\x00', '').replace('\x1e', '-').replace('\x1f', '').replace('\x02', '').replace('\x01', '')
    text = re.sub(r'(?: \| )+\n', '\n', text)
    return _tidy(text.replace('\x0c', '\n\n').replace('\xa0', ' ')), images


def read_doc(data: bytes) -> Reading:
    ole = _ole(data)
    try:
        if not ole.exists('WordDocument'):
            raise ValueError('没有找到文档内容')
        word = ole.openstream('WordDocument').read()
        if struct.unpack_from('<H', word, 0)[0] != 0xA5EC:
            raise ValueError('不是 Word 97 及以上版本的文件')
        if struct.unpack_from('<H', word, 0x0A)[0] & 0x0100:
            raise HTTPException(422, '这份 Word 文件已加密，请先解除密码。')
        table_name = '1Table' if struct.unpack_from('<H', word, 0x0A)[0] & 0x0200 else '0Table'
        table = ole.openstream(table_name).read()
        cslw = struct.unpack_from('<H', word, 0x3E)[0]
        blob = 0x40 + cslw * 4 + 2
        fc_clx, lcb_clx = struct.unpack_from('<II', word, blob + 33 * 8)
        pieces = _doc_pieces(word, table, fc_clx, lcb_clx)
        raw = _doc_text(word, pieces)
        ccp_text, ccp_ftn, ccp_hdd = struct.unpack_from('<III', word, 0x4C)
        ccp_atn, ccp_edn, ccp_txbx = struct.unpack_from('<III', word, 0x5C)
        cursor = ccp_text
        sections = {}
        for name, size in (('脚注', ccp_ftn), ('页眉页脚', ccp_hdd), ('批注', ccp_atn), ('尾注', ccp_edn), ('文本框', ccp_txbx)):
            sections[name] = raw[cursor:cursor + size]
            cursor += size
        main, images = _doc_clean(raw[:ccp_text])
    finally:
        ole.close()
    blocks = [line for line in main.split('\n') if line.strip()]
    for name in ('脚注', '尾注', '文本框'):
        text, _ = _doc_clean(sections.get(name, ''))
        if text.strip():
            blocks.append(f'【{name}】\n' + '\n'.join(f'- {line}' for line in text.split('\n') if line.strip()))
    header_text, _ = _doc_clean(sections.get('页眉页脚', ''))
    header_lines = list(dict.fromkeys(line.strip() for line in header_text.split('\n') if line.strip()))
    if header_lines:
        blocks.append('【页眉页脚】' + ' / '.join(header_lines))
    pages = _split_blocks(blocks)
    _limit_pages(pages, '节')
    notes = [f'文中约有 {images} 处图片或公式对象，里面的文字没有读取。'] if images else []
    return Reading(pages, notes)
