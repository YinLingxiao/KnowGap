"""Small local course library with Chinese/English lexical search and source snapshots."""
import io
import json
import re
from pathlib import PurePosixPath
from typing import Literal
from zipfile import ZipFile

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from fastapi import HTTPException
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.memory.store import Store, ensure_column, new_id, now
from app.models.schemas import Material as BaseMaterial, SourceReference

MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_TEXT = 200_000
MAX_DOCUMENTS = 30
MAX_PAGES = 200
CHUNK_SIZE = 1000
UNIT_LABELS = {'pdf': '页', 'docx': '节', 'pptx': '张幻灯片', 'xlsx': '个工作表', 'text': '页'}


class Material(BaseMaterial):
    """Extended library wire model; preserve the existing fields and legacy text kinds."""
    kind: Literal['pdf', 'txt', 'md', 'csv', 'docx', 'pptx', 'xlsx']


def outline_for(pages: list[tuple[int | None, str]]) -> list[str]:
    headings = [m.group(1).strip() for _, text in pages
                for m in re.finditer(r'^#{1,6}\s+(.+)$', text, re.M)]
    if not headings:
        headings = [text.strip().splitlines()[0] for _, text in pages if text.strip()]
    return list(dict.fromkeys(h[:200] for h in headings))[:20]


def _markdown_row(values) -> str:
    return '| ' + ' | '.join(str(v if v is not None else '').replace('|', '\\|').replace('\n', '<br>')
                             for v in values) + ' |'


def _office_pages(kind: str, data: bytes) -> list[tuple[int, str]]:
    # OOXML is a zip container: bound decompression as well as uploaded bytes.
    with ZipFile(io.BytesIO(data)) as archive:
        if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
            raise HTTPException(413, '文件解压后过大，请按章节拆分。')
    pages = []
    if kind == 'docx':
        document = Document(io.BytesIO(data))
        buffer = ''
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = block.text.strip()
                style = block.style.name if block.style else ''
                heading = re.match(r'(?:Heading|标题)\s*([1-9])$', style, re.I)
                if text and (heading or style == 'Title'):
                    text = '#' * (min(int(heading[1]), 6) if heading else 1) + ' ' + text
            elif isinstance(block, Table):
                rows = [[cell.text for cell in row.cells] for row in block.rows]
                if not any(cell.strip() for row in rows for cell in row):
                    continue
                if rows:
                    rows.insert(1, ['---'] * len(rows[0]))
                text = '\n'.join(_markdown_row(row) for row in rows)
            else:
                continue
            if not text:
                continue
            # Prefer paragraph/table boundaries; split unusually long blocks too.
            for start in range(0, len(text), 1500):
                part = text[start:start + 1500]
                if buffer and len(buffer) + len(part) + 1 > 1500:
                    pages.append((len(pages) + 1, buffer))
                    buffer = ''
                buffer += ('\n' if buffer else '') + part
        if buffer:
            pages.append((len(pages) + 1, buffer))
    elif kind == 'pptx':
        presentation = Presentation(io.BytesIO(data))
        if len(presentation.slides) > MAX_PAGES:
            raise HTTPException(413, f'单份资料最多 {MAX_PAGES} 张幻灯片。')
        for number, slide in enumerate(presentation.slides, 1):
            title = slide.shapes.title
            parts = ['# ' + title.text] if title is not None and title.text.strip() else []
            for shape in slide.shapes:
                if title is not None and shape.shape_id == title.shape_id:
                    continue
                if shape.has_text_frame:
                    parts.append(shape.text)
                if shape.has_table:
                    parts.extend(_markdown_row(cell.text for cell in row.cells) for row in shape.table.rows)
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame
                if notes is not None and notes.text.strip():
                    parts.append('备注：' + notes.text)
            pages.append((number, '\n'.join(parts)))
    else:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if len(workbook.worksheets) > MAX_PAGES:
                raise HTTPException(413, f'单份资料最多 {MAX_PAGES} 个工作表。')
            has_values = False
            for number, sheet in enumerate(workbook.worksheets, 1):
                # Dimensions from the worksheet are the full size, not the preview size.
                if sheet.max_row is None or sheet.max_column is None:
                    sheet.calculate_dimension(force=True)
                rows = list(sheet.iter_rows(max_row=min(sheet.max_row or 0, 200), values_only=True))
                parts = [f'# {sheet.title}', f'总行数：{sheet.max_row}；总列数：{sheet.max_column}；仅展示前 200 行。']
                if not any(v is not None and str(v).strip() for row in rows for v in row):
                    pages.append((number, '\n'.join([*parts, '预览范围内无数据。'])))
                    continue
                has_values = True
                for index, row in enumerate(rows):
                    parts.append(_markdown_row(row))
                    if index == 0:
                        parts.append(_markdown_row('---' for _ in row))
                pages.append((number, '\n'.join(parts)))
            if not has_values:
                raise HTTPException(422, '没有提取到文字，工作表前 200 行内没有可读取的数据。')
        finally:
            workbook.close()
    return pages


def terms(text: str) -> list[str]:
    tokens = []
    for word in re.findall(r"[a-zA-Z][a-zA-Z0-9_]*|[\u3400-\u9fff]+", text.lower()):
        if '\u3400' <= word[0] <= '\u9fff':
            tokens.extend(word)
            tokens.extend(word[i:i + 2] for i in range(len(word) - 1))
        else:
            tokens.append(word)
    return list(dict.fromkeys(tokens))


def chunks(text: str):
    """Keep original line boundaries so TXT/Markdown citations are reproducible."""
    buffer, start, end = '', 1, 1
    for number, line in enumerate(text.splitlines(), 1):
        for offset in range(0, max(len(line), 1), CHUNK_SIZE):
            part = line[offset:offset + CHUNK_SIZE]
            if buffer and len(buffer) + len(part) + 1 > CHUNK_SIZE:
                yield buffer, start, end
                buffer = ''
            if not buffer:
                start = number
            buffer += ('\n' if buffer else '') + part
            end = number
    if buffer.strip():
        yield buffer, start, end


def extract(filename: str, data: bytes) -> tuple[str, list[tuple[int | None, str]]]:
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(413, '单份资料不能超过 20 MB。')
    kind = PurePosixPath(filename).suffix.lower().lstrip('.')
    if kind in {'doc', 'ppt', 'xls'}:
        raise HTTPException(415, '不支持旧二进制格式，请另存为 .docx/.pptx/.xlsx 后上传。')
    if kind not in {'pdf', 'txt', 'md', 'csv', 'docx', 'pptx', 'xlsx'}:
        raise HTTPException(415, '请上传 PDF、DOCX、PPTX、XLSX、TXT、Markdown 或 CSV 文件。')
    if not data:
        raise HTTPException(422, '文件为空，请选择包含文字的资料。')
    if kind == 'pdf':
        if not data.lstrip().startswith(b'%PDF-'):
            raise HTTPException(422, '文件不是有效的 PDF。')
        try:
            reader = PdfReader(io.BytesIO(data), strict=False)
            if reader.is_encrypted:
                raise HTTPException(422, '暂不支持加密 PDF，请先解除密码。')
            if len(reader.pages) > MAX_PAGES:
                raise HTTPException(413, f'单份资料最多 {MAX_PAGES} 页，请按章节拆分。')
            pages, length = [], 0
            for number, page in enumerate(reader.pages, 1):
                text = (page.extract_text() or '').replace('\x00', '')
                length += len(text)
                if length > MAX_TEXT:
                    raise HTTPException(413, '资料文字过多，请按章节拆分。')
                pages.append((number, text))
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, 'PDF 无法解析，请确认文件完整并包含可复制的文字。') from None
    elif kind in {'docx', 'pptx', 'xlsx'}:
        try:
            pages = _office_pages(kind, data)
        except HTTPException:
            raise
        except Exception:
            raise HTTPException(422, '文件无法解析，请确认文件完整、未加密且格式正确。') from None
    else:
        try:
            text = data.decode('utf-8-sig').replace('\x00', '')
        except UnicodeError:
            raise HTTPException(422, '文本资料需使用 UTF-8 编码。') from None
        if len(text) > MAX_TEXT:
            raise HTTPException(413, '资料文字过多，请按章节拆分。')
        pages = [(None, text)]
    if len(pages) > MAX_PAGES:
        raise HTTPException(413, f'单份资料最多 {MAX_PAGES} 页/节/幻灯片/工作表，请按章节拆分。')
    if sum(len(text) for _, text in pages) > MAX_TEXT:
        raise HTTPException(413, '资料文字过多，请按章节拆分。')
    if not any(text.strip() for _, text in pages):
        detail = '这份 PDF 是扫描图片，暂不支持识别文字' if kind == 'pdf' else '没有提取到文字，请选择包含文字的资料。'
        raise HTTPException(422, detail)
    return kind, pages


class MaterialStore:
    def __init__(self, store: Store):
        self.store = store

    def initialize(self):
        with self.store.connection() as db:
            db.executescript('''
                CREATE TABLE IF NOT EXISTS course_materials (
                    id TEXT PRIMARY KEY, student_id TEXT NOT NULL, title TEXT NOT NULL,
                    kind TEXT NOT NULL, page_count INTEGER NOT NULL, chunk_count INTEGER NOT NULL,
                    created_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS course_chunks (
                    id TEXT PRIMARY KEY, document_id TEXT NOT NULL REFERENCES course_materials(id) ON DELETE CASCADE,
                    content TEXT NOT NULL, page INTEGER, start_line INTEGER, end_line INTEGER
                );
                CREATE VIRTUAL TABLE IF NOT EXISTS course_index USING fts5(chunk_id UNINDEXED, terms);
            ''')
            # NULL = not filed under a Subject ("未归类").
            ensure_column(db, 'course_materials', 'subject_id', 'TEXT')
            ensure_column(db, 'course_materials', 'outline_json', "TEXT NOT NULL DEFAULT '[]'")
            ensure_column(db, 'course_materials', 'chars', 'INTEGER NOT NULL DEFAULT 0')

    def add(self, student_id: str, filename: str, data: bytes, subject_id: str | None = None) -> Material:
        if len(data) > MAX_UPLOAD_BYTES:
            raise HTTPException(413, '单份资料不能超过 20 MB。')
        title = PurePosixPath(filename.replace('\\', '/')).name
        title = ''.join(c for c in title if ord(c) >= 32 and ord(c) != 127)[:140] or '课程资料'
        kind, pages = extract(title, data)
        document_id = new_id('material')
        records = [(new_id('source'), document_id, text, page,
                    start if page is None else None, end if page is None else None)
                   for page, content in pages for text, start, end in chunks(content) if text.strip()]
        material = Material(id=document_id, title=title, kind=kind,
                            page_count=len(pages) if kind in {'pdf', 'docx', 'pptx', 'xlsx'} else 0,
                            chunk_count=len(records), created_at=now(), subject_id=subject_id)
        with self.store.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            count = db.execute('SELECT COUNT(*) FROM course_materials WHERE student_id = ?', (student_id,)).fetchone()[0]
            if count >= MAX_DOCUMENTS:
                raise HTTPException(409, '最多保存 30 份资料，请删除不需要的资料后再上传。')
            db.execute('INSERT INTO course_materials(id, student_id, title, kind, page_count, chunk_count, created_at, '
                       'subject_id, outline_json, chars) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)',
                       (document_id, student_id, title, kind, material.page_count, len(records), material.created_at,
                         subject_id, json.dumps(outline_for(pages), ensure_ascii=False), sum(len(text) for _, text in pages)))
            db.executemany('INSERT INTO course_chunks VALUES (?, ?, ?, ?, ?, ?)', records)
            db.executemany('INSERT INTO course_index(chunk_id, terms) VALUES (?, ?)',
                           [(record[0], ' '.join(terms(record[2]))) for record in records])
        return material

    def attachment(self, document_id: str, student_id: str, *, excerpt: bool = False) -> dict | None:
        """Owner-scoped attachment metadata, including a bounded first-turn preview."""
        with self.store.connection() as db:
            row = db.execute('SELECT * FROM course_materials WHERE id = ? AND student_id = ?',
                             (document_id, student_id)).fetchone()
            if row is None:
                return None
            kind = row['kind'] if row['kind'] in UNIT_LABELS else 'text'
            result = dict(id=row['id'], name=row['title'], kind=kind, units=row['page_count'] or 1,
                          unit_label=UNIT_LABELS[kind], chars=row['chars'],
                          outline=json.loads(row['outline_json']), subject_id=row['subject_id'])
            if excerpt:
                # rowid preserves original extraction order, including chunks on the same page.
                preview = ''
                for chunk in db.execute('SELECT content FROM course_chunks WHERE document_id = ? ORDER BY rowid',
                                        (document_id,)):
                    preview += ('\n' if preview else '') + chunk['content']
                    if len(preview) >= 3000:
                        break
                result['excerpt'] = preview[:3000]
            return result

    def list_materials(self, student_id: str, subject_id: str | None = None, unassigned: bool = False) -> list[Material]:
        """All materials by default; one Subject's, or only unassigned ones, when asked."""
        with self.store.connection() as db:
            rows = db.execute('''
                SELECT * FROM course_materials WHERE student_id = ?
                AND (? = 0 OR subject_id IS NULL) AND (? IS NULL OR subject_id = ?) ORDER BY created_at DESC
            ''', (student_id, int(unassigned), subject_id, subject_id)).fetchall()
            return [Material(**dict(row)) for row in rows]

    def delete(self, document_id: str, student_id: str) -> bool:
        with self.store.connection() as db:
            db.execute('BEGIN IMMEDIATE')
            if not db.execute('SELECT id FROM course_materials WHERE id = ? AND student_id = ?', (document_id, student_id)).fetchone():
                return False
            db.execute('DELETE FROM course_index WHERE chunk_id IN (SELECT id FROM course_chunks WHERE document_id = ?)', (document_id,))
            db.execute('DELETE FROM course_materials WHERE id = ?', (document_id,))
            return True

    def search_scoped(self, student_id: str, query: str, limit: int, subject_id: str | None) -> list[SourceReference]:
        """A Subject's own knowledge first, then unassigned materials; other Subjects' files are never mixed in."""
        found = self.search(student_id, query, limit, scope=subject_id) if subject_id else []
        if len(found) < limit:
            found += self.search(student_id, query, limit - len(found), scope=None)
        return [source.model_copy(update={'label': f'资料{index}'}) for index, source in enumerate(found, 1)]

    def search(self, student_id: str, query: str, limit: int = 5, scope: str | None | bool = True,
               document_ids: list[str] | None = None) -> list[SourceReference]:
        """scope: True = every material (legacy behaviour), None = unassigned only, str = that Subject only."""
        tokens = terms(query)
        # Retain only useful single-character terms; bigrams cover normal Chinese text.
        tokens = [t for t in tokens if len(t) > 1 or t in {'栈', '图', '树'}][:64]
        if not tokens or document_ids == []:
            return []
        match = ' OR '.join('"' + token + '"' for token in tokens)
        with self.store.connection() as db:
            document_filter = (' AND d.id IN (' + ','.join('?' for _ in document_ids) + ')') if document_ids else ''
            rows = db.execute('''
                SELECT c.*, d.title FROM course_index i
                JOIN course_chunks c ON c.id = i.chunk_id
                JOIN course_materials d ON d.id = c.document_id
                WHERE course_index MATCH ? AND d.student_id = ?
                AND (? = 1 OR (? IS NULL AND d.subject_id IS NULL) OR d.subject_id = ?)
            ''' + document_filter + ' ORDER BY bm25(course_index), c.id LIMIT ?',
                (match, student_id, int(scope is True), None if scope is True else scope,
                 None if scope is True else scope, *(document_ids or []), limit)).fetchall()
            return [SourceReference(id=row['id'], document_id=row['document_id'], title=row['title'],
                                    label=f'资料{index}', page=row['page'], start_line=row['start_line'],
                                    end_line=row['end_line'], excerpt=row['content'])
                    for index, row in enumerate(rows, 1)]
