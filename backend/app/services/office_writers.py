"""Markdown → .docx / .pptx / .xlsx / .pdf, and in-place edits of .docx / .pptx / .xlsx.

模型只需要写 Markdown：标题、段落、列表、表格、代码块、引用。排版（字体、字号、页边距、表头底色）在这里统一处理，
这样学生拿到的文件在 Word / WPS / PowerPoint 里打开就是整洁可用的，不依赖模型会不会写 python-docx。
"""

import io
import re
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from pptx import Presentation
from pptx.dml.color import RGBColor as PptxColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt as PptPt

SANS = "微软雅黑"
SERIF = "宋体"
MONO = "Consolas"
ACCENT = (0x2F, 0x6F, 0xED)
MAX_TABLE_ROWS = 500
MAX_SLIDES = 60


class DocumentError(ValueError):
    """面向模型和学生的可读错误。"""


# ------------------------------------------------------------------ Markdown 解析

@dataclass
class Block:
    kind: str  # heading | paragraph | bullets | numbered | table | code | quote | break
    text: str = ""
    level: int = 0
    items: list[tuple[int, str]] = field(default_factory=list)  # (缩进层级, 文字)
    rows: list[list[str]] = field(default_factory=list)
    language: str = ""


_HEADING = re.compile(r"^(#{1,6})\s+(.*?)\s*#*\s*$")
_BULLET = re.compile(r"^(\s*)[-*+•]\s+(.*)$")
_NUMBER = re.compile(r"^(\s*)\d+[.)、]\s+(.*)$")
_RULE = re.compile(r"^\s*([-*_])(\s*\1){2,}\s*$")
_TABLE_SEPARATOR = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")


def _cells(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    return [cell.strip().replace("\\|", "|") for cell in re.split(r"(?<!\\)\|", line)]


def parse_markdown(markdown: str) -> list[Block]:
    lines = markdown.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    blocks: list[Block] = []
    paragraph: list[str] = []
    index = 0

    def flush():
        if paragraph:
            blocks.append(Block("paragraph", " ".join(part.strip() for part in paragraph)))
            paragraph.clear()

    while index < len(lines):
        line = lines[index]
        stripped = line.strip()
        if not stripped:
            flush()
        elif stripped.startswith("```"):
            flush()
            language = stripped[3:].strip()
            body = []
            index += 1
            while index < len(lines) and not lines[index].strip().startswith("```"):
                body.append(lines[index])
                index += 1
            blocks.append(Block("code", "\n".join(body), language=language))
        elif _HEADING.match(stripped):
            flush()
            match = _HEADING.match(stripped)
            blocks.append(Block("heading", match.group(2), level=len(match.group(1))))
        elif _RULE.match(stripped):
            flush()
            blocks.append(Block("break"))
        elif "|" in stripped and index + 1 < len(lines) and _TABLE_SEPARATOR.match(lines[index + 1]) and "|" in lines[index + 1]:
            flush()
            rows = [_cells(line)]
            index += 2
            while index < len(lines) and "|" in lines[index] and lines[index].strip():
                rows.append(_cells(lines[index]))
                index += 1
            index -= 1
            width = max(len(row) for row in rows)
            blocks.append(Block("table", rows=[row + [""] * (width - len(row)) for row in rows[:MAX_TABLE_ROWS + 1]]))
        elif _BULLET.match(line) or _NUMBER.match(line):
            flush()
            ordered = bool(_NUMBER.match(line)) and not _BULLET.match(line)
            items: list[tuple[int, str]] = []
            while index < len(lines):
                current = lines[index]
                match = _BULLET.match(current) if not ordered else _NUMBER.match(current)
                if match is None:
                    if current.strip() and current.startswith((" ", "\t")) and items:
                        depth, text = items[-1]
                        items[-1] = (depth, f"{text} {current.strip()}")
                        index += 1
                        continue
                    break
                indent = len(match.group(1).replace("\t", "    "))
                items.append((min(indent // 2, 3), match.group(2).strip()))
                index += 1
            index -= 1
            blocks.append(Block("numbered" if ordered else "bullets", items=items))
        elif stripped.startswith(">"):
            flush()
            quote = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote.append(lines[index].strip().lstrip(">").strip())
                index += 1
            index -= 1
            blocks.append(Block("quote", " ".join(quote)))
        else:
            paragraph.append(line)
        index += 1
    flush()
    return blocks


# 斜体的 * 两侧不能紧挨着西文字母数字（2*3*4 是乘法）；紧挨汉字是正常的
_INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|(?<![\*A-Za-z0-9])\*[^\*\s][^\*]*?\*(?![\*A-Za-z0-9]))")


def inline_runs(text: str) -> list[tuple[str, str]]:
    """[(文字, 样式)]，样式 ∈ plain / bold / italic / code。"""
    runs: list[tuple[str, str]] = []
    position = 0
    for match in _INLINE.finditer(text):
        if match.start() > position:
            runs.append((text[position:match.start()], "plain"))
        token = match.group(0)
        if token.startswith("**"):
            runs.append((token[2:-2], "bold"))
        elif token.startswith("`"):
            runs.append((token[1:-1], "code"))
        else:
            runs.append((token[1:-1], "italic"))
        position = match.end()
    if position < len(text):
        runs.append((text[position:], "plain"))
    return runs or [("", "plain")]


def plain_text(text: str) -> str:
    return "".join(piece for piece, _ in inline_runs(text))


# ------------------------------------------------------------------ Word

def _east_asia(run, font: str) -> None:
    run.font.name = font
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for attribute in ("w:eastAsia", "w:ascii", "w:hAnsi", "w:cs"):
        fonts.set(qn(attribute), font)


def _shade(cell, fill: str) -> None:
    properties = cell._element.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _add_runs(paragraph, text: str, size: float | None = None, color: tuple[int, int, int] | None = None, bold: bool = False):
    for piece, style in inline_runs(text):
        run = paragraph.add_run(piece)
        _east_asia(run, MONO if style == "code" else SANS)
        if size:
            run.font.size = Pt(size)
        run.bold = bold or style == "bold"
        run.italic = style == "italic"
        if style == "code":
            run.font.color.rgb = RGBColor(0xB4, 0x3C, 0x3C)
        elif color:
            run.font.color.rgb = RGBColor(*color)


def _docx_blocks(document, blocks: list[Block]) -> None:
    sizes = {1: 20, 2: 16, 3: 14, 4: 12.5, 5: 12, 6: 12}
    for block in blocks:
        if block.kind == "heading":
            paragraph = document.add_paragraph(style=f"Heading {min(block.level, 4)}")
            _add_runs(paragraph, block.text, size=sizes[block.level], color=(0x1F, 0x2A, 0x44), bold=True)
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = 1.5
            paragraph.paragraph_format.space_after = Pt(6)
            _add_runs(paragraph, block.text, size=11)
        elif block.kind in {"bullets", "numbered"}:
            base = "List Bullet" if block.kind == "bullets" else "List Number"
            for depth, text in block.items:
                style = base if depth == 0 else f"{base} {min(depth + 1, 3)}"
                paragraph = document.add_paragraph(style=style)
                paragraph.paragraph_format.line_spacing = 1.4
                _add_runs(paragraph, text, size=11)
        elif block.kind == "quote":
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.8)
            _add_runs(paragraph, block.text, size=11, color=(0x55, 0x5F, 0x73))
        elif block.kind == "code":
            for line in (block.text or " ").split("\n"):
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.left_indent = Cm(0.4)
                paragraph.paragraph_format.space_after = Pt(0)
                run = paragraph.add_run(line or " ")
                _east_asia(run, MONO)
                run.font.size = Pt(9.5)
                shading = OxmlElement("w:shd")
                shading.set(qn("w:val"), "clear")
                shading.set(qn("w:color"), "auto")
                shading.set(qn("w:fill"), "F3F5F8")
                paragraph._element.get_or_add_pPr().append(shading)
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "table":
            width = len(block.rows[0])
            table = document.add_table(rows=len(block.rows), cols=width)
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            for r, row in enumerate(block.rows):
                for c, text in enumerate(row):
                    cell = table.cell(r, c)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    _add_runs(paragraph, text, size=10, bold=r == 0)
                    if r == 0:
                        _shade(cell, "E8EEFB")
            document.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "break":
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_docx(title: str, markdown: str) -> bytes:
    blocks = parse_markdown(markdown)
    document = Document()
    section = document.sections[0]
    section.left_margin = section.right_margin = Cm(2.6)
    section.top_margin = section.bottom_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = SANS
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), SANS)
    document.core_properties.title = title
    # 内容自带一级标题时不再重复写文件名
    if not (blocks and blocks[0].kind == "heading" and blocks[0].level == 1):
        head = document.add_paragraph(style="Title")
        _add_runs(head, title, size=24, color=(0x1F, 0x2A, 0x44), bold=True)
    _docx_blocks(document, blocks)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ PowerPoint

def _slide_sections(blocks: list[Block]) -> list[tuple[str, list[Block], str]]:
    """一个 `##` 一页；`#` 只作为整份演示的标题。`备注：` 开头的段落进入演讲者备注。"""
    slides: list[tuple[str, list[Block], str]] = []
    title, body, notes = None, [], []

    def close():
        if title is not None:
            slides.append((title, body[:], "\n".join(notes)))

    for block in blocks:
        if block.kind == "heading" and block.level >= 2:
            close()
            title, body, notes = block.text, [], []
        elif title is None:
            continue
        elif block.kind == "paragraph" and re.match(r"^(备注|讲稿|Notes?)[：:]", block.text):
            notes.append(re.sub(r"^(备注|讲稿|Notes?)[：:]\s*", "", block.text))
        else:
            body.append(block)
    close()
    return slides


def _ppt_text(frame, text: str, size: float, color=None, bold=False, bullet_level: int | None = None):
    paragraph = frame.add_paragraph() if frame.paragraphs[0].runs else frame.paragraphs[0]
    paragraph.level = bullet_level or 0
    for piece, style in inline_runs(text):
        run = paragraph.add_run()
        run.text = piece
        run.font.name = MONO if style == "code" else SANS
        run.font.size = PptPt(size)
        run.font.bold = bold or style == "bold"
        run.font.italic = style == "italic"
        if color is not None:
            run.font.color.rgb = PptxColor(*color)
    return paragraph


def _fit_size(count: int, longest: int) -> float:
    if count <= 4 and longest <= 28:
        return 24
    if count <= 6 and longest <= 40:
        return 20
    if count <= 8:
        return 18
    return 15


def build_pptx(title: str, markdown: str, subtitle: str = "") -> bytes:
    blocks = parse_markdown(markdown)
    slides = _slide_sections(blocks)
    if not slides:
        raise DocumentError("演示文稿内容里没有找到 `## 页标题`，每一页要用二级标题开头。")
    if len(slides) > MAX_SLIDES:
        raise DocumentError(f"一份演示最多 {MAX_SLIDES} 页，请拆成几份。")
    deck_title = next((block.text for block in blocks if block.kind == "heading" and block.level == 1), title)
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = Inches(13.333), Inches(7.5)
    width, height = presentation.slide_width, presentation.slide_height

    cover = presentation.slides.add_slide(presentation.slide_layouts[6])
    band = cover.shapes.add_shape(1, 0, 0, Inches(0.45), height)
    band.fill.solid(); band.fill.fore_color.rgb = PptxColor(*ACCENT); band.line.fill.background()
    box = cover.shapes.add_textbox(Inches(1.2), Inches(2.5), width - Inches(2.4), Inches(1.6))
    box.text_frame.word_wrap = True
    _ppt_text(box.text_frame, deck_title, 44, color=(0x1F, 0x2A, 0x44), bold=True)
    if subtitle:
        sub = cover.shapes.add_textbox(Inches(1.2), Inches(4.2), width - Inches(2.4), Inches(0.9))
        sub.text_frame.word_wrap = True
        _ppt_text(sub.text_frame, subtitle, 20, color=(0x55, 0x5F, 0x73))

    for number, (slide_title, body, notes) in enumerate(slides, start=1):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        _slide_content(slide, width, height, slide_title, body, number)
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue()


def _slide_content(slide, width, height, slide_title: str, body: list[Block], number: int | None = None) -> None:
    bar = slide.shapes.add_shape(1, 0, 0, Inches(0.18), height)
    bar.fill.solid(); bar.fill.fore_color.rgb = PptxColor(*ACCENT); bar.line.fill.background()
    head = slide.shapes.add_textbox(Inches(0.7), Inches(0.4), width - Inches(1.4), Inches(1.0))
    head.text_frame.word_wrap = True
    _ppt_text(head.text_frame, slide_title, 32, color=(0x1F, 0x2A, 0x44), bold=True)
    top, bottom = Inches(1.6), height - Inches(0.7)
    area_width = width - Inches(1.4)

    tables = [block for block in body if block.kind == "table"]
    texts = [block for block in body if block.kind != "table"]
    items: list[tuple[int, str]] = []
    for block in texts:
        if block.kind in {"bullets", "numbered"}:
            items.extend(block.items)
        elif block.kind == "code":
            items.extend((1, line) for line in block.text.splitlines() if line.strip())
        elif block.kind in {"paragraph", "quote", "heading"}:
            items.append((0, block.text))
    text_height = bottom - top if not tables else int((bottom - top) * 0.45)
    if items:
        box = slide.shapes.add_textbox(Inches(0.7), top, area_width, text_height)
        frame = box.text_frame
        frame.word_wrap = True
        size = _fit_size(len(items), max(len(text) for _, text in items))
        for depth, text in items:
            paragraph = _ppt_text(frame, ("•  " if depth == 0 else "–  ") + text, size - 2 * depth, color=(0x25, 0x2B, 0x3A))
            paragraph.space_after = PptPt(8)
            if depth:
                paragraph.level = min(depth, 4)
    for table_block in tables[:1]:
        rows, columns = len(table_block.rows), len(table_block.rows[0])
        top_table = top + (text_height + Inches(0.2) if items else 0)
        shape = slide.shapes.add_table(rows, columns, Inches(0.7), top_table, area_width,
                                       min(Inches(0.5) * rows, bottom - top_table))
        for r, row in enumerate(table_block.rows):
            for c, text in enumerate(row):
                cell = shape.table.cell(r, c)
                cell.text = ""
                _ppt_text(cell.text_frame, text, 14 if rows < 8 else 11, bold=r == 0,
                          color=(0xFF, 0xFF, 0xFF) if r == 0 else (0x25, 0x2B, 0x3A))
                if r == 0:
                    cell.fill.solid(); cell.fill.fore_color.rgb = PptxColor(*ACCENT)
    if number is not None:
        footer = slide.shapes.add_textbox(width - Inches(1.4), height - Inches(0.55), Inches(0.9), Inches(0.35))
        paragraph = _ppt_text(footer.text_frame, str(number), 12, color=(0x88, 0x92, 0xA6))
        paragraph.alignment = PP_ALIGN.RIGHT


# ------------------------------------------------------------------ Excel

def _number(text: str):
    cleaned = text.strip().replace(",", "") if re.fullmatch(r"-?[\d,]+(\.\d+)?", text.strip()) else text.strip()
    if re.fullmatch(r"-?\d+", cleaned):
        return int(cleaned)
    if re.fullmatch(r"-?\d+\.\d+", cleaned):
        return float(cleaned)
    if cleaned.endswith("%") and re.fullmatch(r"-?\d+(\.\d+)?%", cleaned):
        return float(cleaned[:-1]) / 100
    return text


def build_xlsx(title: str, markdown: str) -> bytes:
    blocks = parse_markdown(markdown)
    workbook = Workbook()
    workbook.remove(workbook.active)
    name, used = "Sheet1", set()
    pending: list[tuple[str, list[list[str]]]] = []
    for block in blocks:
        if block.kind == "heading":
            name = plain_text(block.text)
        elif block.kind == "table":
            pending.append((name, block.rows))
    if not pending:
        raise DocumentError("表格内容里没有找到 Markdown 表格（| 列 | 列 |），请用表格给出数据。")
    thin = Side(style="thin", color="D0D5DD")
    for sheet_name, rows in pending:
        clean = re.sub(r"[\\/*?:\[\]]", " ", sheet_name).strip()[:28] or "Sheet"
        candidate, suffix = clean, 2
        while candidate in used:
            candidate, suffix = f"{clean[:25]}_{suffix}", suffix + 1
        used.add(candidate)
        sheet = workbook.create_sheet(candidate)
        for r, row in enumerate(rows, start=1):
            for c, text in enumerate(row, start=1):
                value = plain_text(text)
                cell = sheet.cell(row=r, column=c, value=value if value.startswith("=") else _number(value))
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if r == 1:
                    cell.font = Font(name=SANS, bold=True, color="FFFFFF")
                    cell.fill = PatternFill("solid", fgColor="2F6FED")
                else:
                    cell.font = Font(name=SANS)
        for c in range(1, len(rows[0]) + 1):
            longest = max((sum(2 if ord(ch) > 255 else 1 for ch in str(sheet.cell(row=r, column=c).value or "")) for r in range(1, len(rows) + 1)), default=8)
            sheet.column_dimensions[get_column_letter(c)].width = min(max(longest + 3, 8), 48)
        sheet.freeze_panes = "A2"
    out = io.BytesIO()
    workbook.properties.title = title
    workbook.save(out)
    return out.getvalue()


# ------------------------------------------------------------------ PDF

_PDF_FONTS = [
    ("msyh", r"C:\Windows\Fonts\msyh.ttc", 0), ("simsun", r"C:\Windows\Fonts\simsun.ttc", 0),
    ("simhei", r"C:\Windows\Fonts\simhei.ttf", None),
    ("noto", "/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc", 0),
    ("pingfang", "/System/Library/Fonts/PingFang.ttc", 0),
]


def _pdf_font() -> tuple[str, str]:
    """(正文字体, 等宽字体)。优先嵌入系统里的中文字体；都没有时用 PDF 内置的宋体 CID 字体。"""
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.cidfonts import UnicodeCIDFont
    from reportlab.pdfbase.ttfonts import TTFont

    for name, path, index in _PDF_FONTS:
        if Path(path).is_file():
            try:
                if name not in pdfmetrics.getRegisteredFontNames():
                    pdfmetrics.registerFont(TTFont(name, path, subfontIndex=index or 0))
                return name, name
            except Exception:
                continue
    if "STSong-Light" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    return "STSong-Light", "STSong-Light"


def _escape(text: str) -> str:
    return text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _pdf_inline(text: str, mono: str) -> str:
    parts = []
    for piece, style in inline_runs(text):
        piece = _escape(piece)
        if style == "bold":
            piece = f"<b>{piece}</b>"
        elif style == "italic":
            piece = f"<i>{piece}</i>"
        elif style == "code":
            piece = f'<font name="{mono}" color="#B43C3C">{piece}</font>'
        parts.append(piece)
    return "".join(parts)


def build_pdf(title: str, markdown: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        HRFlowable, ListFlowable, ListItem, PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer, Table, TableStyle,
    )

    blocks = parse_markdown(markdown)
    font, mono = _pdf_font()

    def style(name, size, leading=None, **extra):
        options = dict(fontName=font, fontSize=size, leading=leading or size * 1.6, alignment=TA_LEFT,
                       wordWrap="CJK", textColor=colors.HexColor("#252B3A"))
        options.update(extra)
        return ParagraphStyle(name, **options)

    body = style("body", 10.5, spaceAfter=6)
    headings = {n: style(f"h{n}", size, size * 1.4, spaceBefore=10, spaceAfter=6, textColor=colors.HexColor("#1F2A44"))
                for n, size in {1: 20, 2: 16, 3: 13.5, 4: 12}.items()}
    quote = style("quote", 10.5, leftIndent=14, textColor=colors.HexColor("#555F73"), spaceAfter=6)
    code = ParagraphStyle("code", fontName=mono, fontSize=8.8, leading=12, backColor=colors.HexColor("#F3F5F8"),
                          leftIndent=6, rightIndent=6, borderPadding=4, spaceAfter=8, wordWrap="CJK")
    cell = style("cell", 9, 13)
    head_cell = style("head", 9, 13, textColor=colors.white)

    story = []
    if not (blocks and blocks[0].kind == "heading" and blocks[0].level == 1):
        story.append(Paragraph(_pdf_inline(title, mono), headings[1]))
    for block in blocks:
        if block.kind == "heading":
            story.append(Paragraph(_pdf_inline(block.text, mono), headings[min(block.level, 4)]))
        elif block.kind == "paragraph":
            story.append(Paragraph(_pdf_inline(block.text, mono), body))
        elif block.kind in {"bullets", "numbered"}:
            story.append(ListFlowable(
                [ListItem(Paragraph(_pdf_inline(text, mono), body), leftIndent=14 + depth * 12) for depth, text in block.items],
                bulletType="bullet" if block.kind == "bullets" else "1", start="•" if block.kind == "bullets" else 1,
                bulletFontName=font, leftIndent=14))
        elif block.kind == "quote":
            story.append(Paragraph(_pdf_inline(block.text, mono), quote))
        elif block.kind == "code":
            story.append(Preformatted(block.text or " ", code))
        elif block.kind == "table":
            data = [[Paragraph(_pdf_inline(text, mono), head_cell if r == 0 else cell) for text in row]
                    for r, row in enumerate(block.rows)]
            table = Table(data, repeatRows=1, hAlign="LEFT")
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2F6FED")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#D0D5DD")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F9FC")]),
            ]))
            story += [table, Spacer(1, 8)]
        elif block.kind == "break":
            story.append(PageBreak())
    if not story:
        raise DocumentError("文档内容是空的。")
    out = io.BytesIO()
    SimpleDocTemplate(out, pagesize=A4, leftMargin=2.2 * cm, rightMargin=2.2 * cm, topMargin=2 * cm, bottomMargin=2 * cm,
                      title=title).build(story)
    return out.getvalue()


# ------------------------------------------------------------------ 编辑

def _replace_in_paragraph(paragraph, find: str, replacement: str, only_first: bool) -> int:
    """跨 run 替换：保留起点所在 run 的格式，被替换范围内其余 run 的相应文字删掉。

    替换后从替换文字之后继续找，所以 replacement 里含有 find（如「A」→「A（说明）」）也不会无限循环。
    """
    count, cursor = 0, 0
    while True:
        runs = paragraph.runs
        text = "".join(run.text for run in runs)
        start = text.find(find, cursor)
        if start < 0 or (only_first and count):
            return count
        end, position, first = start + len(find), 0, True
        for run in runs:
            run_start, run_end = position, position + len(run.text)
            position = run_end
            if run_end <= start or run_start >= end:
                continue
            local_start, local_end = max(start, run_start) - run_start, min(end, run_end) - run_start
            run.text = run.text[:local_start] + (replacement if first else "") + run.text[local_end:]
            first = False
        count += 1
        cursor = start + len(replacement)
        if only_first:
            return count


def _docx_paragraphs(document):
    for paragraph in document.paragraphs:
        yield paragraph
    tables = list(document.tables)
    while tables:
        table = tables.pop()
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    yield paragraph
                tables.extend(cell.tables)
    for section in document.sections:
        for part in (section.header, section.footer):
            for paragraph in part.paragraphs:
                yield paragraph


def edit_docx(data: bytes, operations: list[dict]) -> tuple[bytes, list[str]]:
    document = Document(io.BytesIO(data))
    log: list[str] = []
    for number, operation in enumerate(operations, start=1):
        kind = operation.get("op")
        if kind == "replace":
            find, replacement = str(operation.get("find") or ""), str(operation.get("replace") or "")
            if not find:
                raise DocumentError(f"第 {number} 项操作缺少 find。")
            total = sum(_replace_in_paragraph(p, find, replacement, not operation.get("all", True)) for p in _docx_paragraphs(document))
            if total == 0:
                raise DocumentError(f"第 {number} 项操作：文档里没有找到「{find[:30]}」。请先读取文档确认原文。")
            log.append(f"替换「{find[:20]}」{total} 处")
        elif kind == "append":
            blocks = parse_markdown(str(operation.get("markdown") or ""))
            if not blocks:
                raise DocumentError(f"第 {number} 项操作的 markdown 是空的。")
            _docx_blocks(document, blocks)
            log.append("在末尾追加内容")
        elif kind == "insert_after":
            anchor = str(operation.get("after") or "")
            target = next((p for p in document.paragraphs if anchor and anchor in p.text), None)
            if target is None:
                raise DocumentError(f"第 {number} 项操作：没有找到包含「{anchor[:30]}」的段落。")
            scratch = Document()
            _docx_blocks(scratch, parse_markdown(str(operation.get("markdown") or "")))
            anchor_element = target._element
            for element in list(scratch.element.body)[:-1] if scratch.element.body[-1].tag.endswith("sectPr") else list(scratch.element.body):
                anchor_element.addnext(element)
                anchor_element = element
            log.append(f"在「{anchor[:20]}」之后插入内容")
        elif kind == "delete":
            find = str(operation.get("find") or "")
            targets = [p for p in document.paragraphs if find and find in p.text]
            if not targets:
                raise DocumentError(f"第 {number} 项操作：没有找到包含「{find[:30]}」的段落。")
            for paragraph in targets:
                paragraph._element.getparent().remove(paragraph._element)
            log.append(f"删除 {len(targets)} 个包含「{find[:20]}」的段落")
        else:
            raise DocumentError(f"第 {number} 项操作的 op 不支持：{kind!r}（可用 replace / append / insert_after / delete）。")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue(), log


def _pptx_frames(presentation):
    for slide in presentation.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                yield shape.text_frame
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    for cell in row.cells:
                        yield cell.text_frame
        if slide.has_notes_slide:
            yield slide.notes_slide.notes_text_frame


def edit_pptx(data: bytes, operations: list[dict]) -> tuple[bytes, list[str]]:
    presentation = Presentation(io.BytesIO(data))
    log: list[str] = []
    for number, operation in enumerate(operations, start=1):
        kind = operation.get("op")
        if kind == "replace":
            find, replacement = str(operation.get("find") or ""), str(operation.get("replace") or "")
            if not find:
                raise DocumentError(f"第 {number} 项操作缺少 find。")
            total = sum(_replace_in_paragraph(p, find, replacement, False) for frame in _pptx_frames(presentation) for p in frame.paragraphs)
            if total == 0:
                raise DocumentError(f"第 {number} 项操作：幻灯片里没有找到「{find[:30]}」。")
            log.append(f"替换「{find[:20]}」{total} 处")
        elif kind == "append_slides":
            slides = _slide_sections(parse_markdown(str(operation.get("markdown") or "")))
            if not slides:
                raise DocumentError(f"第 {number} 项操作：markdown 里没有 `## 页标题`。")
            if len(presentation.slides) + len(slides) > MAX_SLIDES:
                raise DocumentError(f"追加后超过 {MAX_SLIDES} 页。")
            blank = next((layout for layout in presentation.slide_layouts if layout.name.lower() == "blank"), presentation.slide_layouts[-1])
            for slide_title, body, notes in slides:
                slide = presentation.slides.add_slide(blank)
                _slide_content(slide, presentation.slide_width, presentation.slide_height, slide_title, body, len(presentation.slides))
                if notes:
                    slide.notes_slide.notes_text_frame.text = notes
            log.append(f"追加 {len(slides)} 页")
        elif kind == "delete_slide":
            index = int(operation.get("slide") or 0)
            if not 1 <= index <= len(presentation.slides):
                raise DocumentError(f"第 {number} 项操作：没有第 {index} 页（共 {len(presentation.slides)} 页）。")
            slide_ids = presentation.slides._sldIdLst
            presentation.part.drop_rel(slide_ids[index - 1].rId)
            del slide_ids[index - 1]
            log.append(f"删除第 {index} 页")
        elif kind == "set_notes":
            index = int(operation.get("slide") or 0)
            if not 1 <= index <= len(presentation.slides):
                raise DocumentError(f"第 {number} 项操作：没有第 {index} 页。")
            presentation.slides[index - 1].notes_slide.notes_text_frame.text = str(operation.get("text") or "")
            log.append(f"更新第 {index} 页备注")
        else:
            raise DocumentError(f"第 {number} 项操作的 op 不支持：{kind!r}（可用 replace / append_slides / delete_slide / set_notes）。")
    out = io.BytesIO()
    presentation.save(out)
    return out.getvalue(), log


def edit_xlsx(data: bytes, operations: list[dict]) -> tuple[bytes, list[str]]:
    workbook = load_workbook(io.BytesIO(data))
    log: list[str] = []

    def sheet_of(operation, number):
        name = operation.get("sheet")
        if not name:
            return workbook.active
        if name not in workbook.sheetnames:
            raise DocumentError(f"第 {number} 项操作：没有名为「{name}」的工作表（现有：{'、'.join(workbook.sheetnames)}）。")
        return workbook[name]

    for number, operation in enumerate(operations, start=1):
        kind = operation.get("op")
        if kind == "set_cell":
            cell = str(operation.get("cell") or "")
            if not re.fullmatch(r"[A-Za-z]{1,3}[1-9]\d{0,5}", cell):
                raise DocumentError(f"第 {number} 项操作：单元格地址 {cell!r} 无效。")
            value = operation.get("value")
            sheet_of(operation, number)[cell.upper()] = _number(value) if isinstance(value, str) and not value.startswith("=") else value
            log.append(f"写入 {cell.upper()}")
        elif kind == "append_rows":
            rows = operation.get("rows")
            if not isinstance(rows, list) or not rows or not all(isinstance(row, list) for row in rows):
                raise DocumentError(f"第 {number} 项操作：rows 必须是二维数组。")
            sheet = sheet_of(operation, number)
            for row in rows[:MAX_TABLE_ROWS]:
                sheet.append([_number(item) if isinstance(item, str) and not item.startswith("=") else item for item in row])
            log.append(f"追加 {min(len(rows), MAX_TABLE_ROWS)} 行")
        elif kind == "replace":
            find, replacement = str(operation.get("find") or ""), str(operation.get("replace") or "")
            total = 0
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if isinstance(cell.value, str) and find and find in cell.value and not cell.value.startswith("="):
                            cell.value = cell.value.replace(find, replacement)
                            total += 1
            if total == 0:
                raise DocumentError(f"第 {number} 项操作：表格里没有找到「{find[:30]}」。")
            log.append(f"替换「{find[:20]}」{total} 处")
        else:
            raise DocumentError(f"第 {number} 项操作的 op 不支持：{kind!r}（可用 set_cell / append_rows / replace）。")
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue(), log


BUILDERS = {"docx": build_docx, "pdf": build_pdf, "xlsx": build_xlsx}
EDITORS = {"docx": edit_docx, "pptx": edit_pptx, "xlsx": edit_xlsx}
MIME = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
    "md": "text/markdown; charset=utf-8",
}
