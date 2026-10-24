"""课件里的图片：按幻灯片取出，交给「服务模型分工」里的识图模型转写，文字并回对应的那一页。

只读文字的话，截图、位图、流程图里的内容学生以为 AI 看到了，其实没有。这里补上这一步：
取图 → 去掉装饰图（太小、同一张在很多页重复出现的 logo）→ 缩放 → 识图 → 以「[图片内容]」附在该页文字后面。
没配置识图模型、或某张识别失败，都不影响已经读到的文字，并如实写进读取说明。
"""

import asyncio
import hashlib
import io
import struct
import time
from dataclasses import dataclass

import httpx

from app.agents.llm import ProviderError
from app.services.office_readers import (
    Reading, _RT_SLIDE, _ole, _persist_directory, _ppt_slide_containers, _records,
)
from app.services.pdf_ocr import OCR_CONCURRENCY, RENDER_EDGE, _ocr_page

MAX_PICTURES = 24            # 一份课件最多识别的图片数，控制上传耗时和调用量
MIN_EDGE = 80                # 短边小于它的是图标、项目符号
MIN_AREA = 120 * 120
DECORATION_SLIDES = 3        # 同一张图出现在超过这么多页，当作 logo / 页脚装饰
DEADLINE_SECONDS = 150       # 整份课件的识图总时长，超过后不再新开识别
SLIDE_PROMPT = (
    "这是课件某一页幻灯片里的一张图片。请把图中的信息转写成文字，供学生检索和提问：\n"
    "1. 图里有文字、数字、代码、地址、寄存器或表格，逐字转写，保持阅读顺序；表格用 Markdown 表格。\n"
    "2. 流程图、结构图、内存/寄存器分配图、电路图、示意图：用简洁的文字列出各个部分、标注、数值，以及它们之间的位置或连接关系。\n"
    "3. 只是装饰、背景、图标、没有信息量的图，什么都不要输出。\n"
    "看不清的字写「□」，不要猜测、补写或总结。只输出内容，不要加任何说明。"
)

_RT_PPDRAWING_GROUP, _RT_PPDRAWING = 1035, 1036
_ESCHER_DGG, _ESCHER_BSTORE, _ESCHER_FBSE = 0xF000, 0xF001, 0xF007
_ESCHER_OPT = (0xF00B, 0xF121, 0xF122)
_PIB = 0x0104                # 图片填充/图片形状引用的 BLIP 序号（从 1 开始）
_BLIP_JPEG, _BLIP_PNG, _BLIP_DIB, _BLIP_JPEG_CMYK = 0xF01D, 0xF01E, 0xF01F, 0xF02A


@dataclass
class Picture:
    slide: int
    data: bytes


@dataclass
class Collected:
    pictures: list[Picture]
    total: int = 0            # 课件里找到的图片总数（含装饰图）
    decorative: int = 0
    unsupported: int = 0      # 矢量图（EMF/WMF）等没法转成位图的
    over_limit: int = 0


# ------------------------------------------------------------------ 取图

def _pptx_raw(data: bytes) -> list[tuple[int, bytes]]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    found: list[tuple[int, bytes]] = []

    def walk(shapes, number: int):
        for shape in shapes:
            if getattr(shape, 'shape_type', None) == MSO_SHAPE_TYPE.GROUP:
                walk(shape.shapes, number)
                continue
            try:
                blob = shape.image.blob
            except Exception:  # noqa: BLE001 - 不是图片形状
                continue
            found.append((number, blob))

    for number, slide in enumerate(Presentation(io.BytesIO(data)).slides, 1):
        walk(slide.shapes, number)
    return found


def _escher(buf: bytes, start: int, end: int):
    """Office Art 记录的头部和 PPT 记录一致，复用同一个读取器。"""
    return _records(buf, start, end)


def _bstore_entries(buf: bytes, pictures: bytes) -> list[tuple[int, bytes] | None]:
    """BLIP 序号（0 起）→ (类型, 图片字节)。读不出的位置是 None，序号保持对应。"""
    entries: list[tuple[int, bytes] | None] = []
    for rtype, _v, _i, body, length in _records(buf, 0, len(buf)):
        if rtype != 1000:  # Document
            continue
        for rtype2, _v2, _i2, body2, length2 in _records(buf, body, body + length):
            if rtype2 != _RT_PPDRAWING_GROUP:
                continue
            for rtype3, _v3, _i3, body3, length3 in _escher(buf, body2, body2 + length2):
                if rtype3 != _ESCHER_DGG:
                    continue
                for rtype4, _v4, _i4, body4, length4 in _escher(buf, body3, body3 + length3):
                    if rtype4 != _ESCHER_BSTORE:
                        continue
                    for rtype5, _v5, _i5, body5, length5 in _escher(buf, body4, body4 + length4):
                        entries.append(_fbse_blip(buf, body5, length5, pictures) if rtype5 == _ESCHER_FBSE else None)
    return entries


def _fbse_blip(buf: bytes, body: int, length: int, pictures: bytes) -> tuple[int, bytes] | None:
    if length < 36:
        return None
    name_length = buf[body + 33]
    embedded = body + 36 + name_length
    if embedded + 8 <= body + length:  # BLIP 直接跟在 FBSE 后面
        return _blip(buf, embedded)
    offset = struct.unpack_from('<I', buf, body + 28)[0]
    if offset == 0xFFFFFFFF or offset + 8 > len(pictures):
        return None
    return _blip(pictures, offset)


def _blip(buf: bytes, offset: int) -> tuple[int, bytes] | None:
    ver_inst, rtype, length = struct.unpack_from('<HHI', buf, offset)
    if rtype not in (_BLIP_JPEG, _BLIP_PNG, _BLIP_DIB, _BLIP_JPEG_CMYK):
        return None  # EMF / WMF / PICT / TIFF：没有可靠的位图转换
    instance = ver_inst >> 4
    two_uids = instance in (0x46B, 0x6E3, 0x6E1, 0x7A9)
    start = offset + 8 + (32 if two_uids else 16) + 1
    end = offset + 8 + length
    if start >= end or end > len(buf):
        return None
    payload = buf[start:end]
    if rtype == _BLIP_DIB:
        payload = _dib_to_bmp(payload)
    return rtype, payload


def _dib_to_bmp(dib: bytes) -> bytes:
    header = struct.unpack_from('<I', dib, 0)[0] if len(dib) >= 4 else 40
    bits, used = (struct.unpack_from('<H', dib, 14)[0], struct.unpack_from('<I', dib, 32)[0]) if len(dib) >= 36 else (24, 0)
    palette = (used or (1 << bits)) * 4 if bits <= 8 else 0
    return b'BM' + struct.pack('<IHHI', 14 + len(dib), 0, 0, 14 + header + palette) + dib


def _pib_indexes(buf: bytes, start: int, end: int) -> list[int]:
    """一张幻灯片的绘图里，按先后顺序引用到的 BLIP 序号（从 1 起）。"""
    found: list[int] = []
    for rtype, ver, inst, body, length in _escher(buf, start, end):
        if ver == 0xF:
            found.extend(_pib_indexes(buf, body, body + length))
        elif rtype in _ESCHER_OPT:
            for index in range(inst):
                position = body + index * 6
                if position + 6 > body + length:
                    break
                opid, value = struct.unpack_from('<Hi', buf, position)
                if opid & 0x3FFF == _PIB and not opid & 0x8000 and value > 0:
                    found.append(value)
    return found


def _ppt_raw(data: bytes) -> list[tuple[int, bytes]]:
    ole = _ole(data)
    try:
        if not ole.exists('PowerPoint Document') or not ole.exists('Pictures'):
            return []
        buf = ole.openstream('PowerPoint Document').read()
        pictures = ole.openstream('Pictures').read()
        directory = _persist_directory(ole, buf)
        containers, _notes = _ppt_slide_containers(buf, directory)
    finally:
        ole.close()
    blips = _bstore_entries(buf, pictures)
    found: list[tuple[int, bytes]] = []
    for number, offset in enumerate(containers, 1):
        ver_inst, rtype, length = struct.unpack_from('<HHI', buf, offset)
        if rtype != _RT_SLIDE:
            continue
        for rtype2, _v, _i, body2, length2 in _records(buf, offset + 8, offset + 8 + length):
            if rtype2 != _RT_PPDRAWING:
                continue
            seen: set[int] = set()
            for index in _pib_indexes(buf, body2, body2 + length2):
                if index in seen or index > len(blips) or blips[index - 1] is None:
                    continue
                seen.add(index)
                found.append((number, blips[index - 1][1]))
    return found


# ------------------------------------------------------------------ 过滤与缩放

def _to_png(raw: bytes) -> tuple[bytes, int, int] | None:
    from PIL import Image

    try:
        image = Image.open(io.BytesIO(raw))
        image.load()
    except Exception:  # noqa: BLE001 - 矢量图、损坏的图
        return None
    width, height = image.size
    if image.mode in ('RGBA', 'LA', 'P'):
        image = image.convert('RGBA')
        canvas = Image.new('RGB', image.size, (255, 255, 255))
        canvas.paste(image, mask=image.getchannel('A'))
        image = canvas
    else:
        image = image.convert('RGB')
    edge = max(image.size)
    if edge > RENDER_EDGE:
        scale = RENDER_EDGE / edge
        image = image.resize((max(1, round(image.width * scale)), max(1, round(image.height * scale))))
    buffer = io.BytesIO()
    image.save(buffer, format='PNG', optimize=True)
    return buffer.getvalue(), width, height


def collect_pictures(kind: str, data: bytes) -> Collected:
    raw = _pptx_raw(data) if data.startswith(b'PK') else _ppt_raw(data)
    result = Collected([], total=len(raw))
    slides_of: dict[str, set[int]] = {}
    for number, blob in raw:
        slides_of.setdefault(hashlib.sha1(blob).hexdigest(), set()).add(number)
    taken: set[tuple[int, str]] = set()
    for number, blob in raw:
        digest = hashlib.sha1(blob).hexdigest()
        if len(slides_of[digest]) > DECORATION_SLIDES or (number, digest) in taken:
            result.decorative += 1
            continue
        converted = _to_png(blob)
        if converted is None:
            result.unsupported += 1
            continue
        png, width, height = converted
        if min(width, height) < MIN_EDGE or width * height < MIN_AREA:
            result.decorative += 1
            continue
        if len(result.pictures) >= MAX_PICTURES:
            result.over_limit += 1
            continue
        taken.add((number, digest))
        result.pictures.append(Picture(number, png))
    return result


# ------------------------------------------------------------------ 识图并回

def _drop_old_notes(notes: list[str]) -> list[str]:
    return [note for note in notes if '图片里的文字没有读取' not in note and '图片、图表里的文字没有读取' not in note]


def _numbers(values) -> str:
    ordered = sorted(set(values))
    return '、'.join(str(value) for value in ordered[:12]) + (' 等' if len(ordered) > 12 else '')


async def add_picture_text(kind: str, data: bytes, reading: Reading, service_models, *,
                           transport: httpx.AsyncBaseTransport | None = None) -> Reading:
    """返回并入图片文字后的 Reading；没有图片或识图不可用时返回带说明的原内容，不抛错。"""
    found = await asyncio.to_thread(collect_pictures, kind, data)
    if not found.pictures:
        if found.unsupported:
            notes = _drop_old_notes(reading.notes) + [f'课件里有 {found.unsupported} 张矢量图（EMF/WMF）没有识别，图里的文字没有读取。']
            return Reading(reading.pages, notes)
        return reading
    config = service_models.vision_config(service_models.config())
    if not config.configured:
        notes = _drop_old_notes(reading.notes) + [
            f'课件里有 {len(found.pictures)} 张内容图片，图片里的文字没有读取。'
            '到「设置 → 模型接入」配置一个多模态识图模型后重新上传，就能识别。']
        return Reading(reading.pages, notes)

    gate = asyncio.Semaphore(OCR_CONCURRENCY)
    deadline = time.monotonic() + DEADLINE_SECONDS

    async def one(picture: Picture):
        async with gate:
            if time.monotonic() > deadline:
                return picture.slide, '', '超时'
            try:
                return picture.slide, await _ocr_page(picture.data, config, transport, SLIDE_PROMPT), ''
            except ProviderError as exc:
                return picture.slide, '', exc.detail

    results = await asyncio.gather(*(one(picture) for picture in found.pictures))
    by_slide: dict[int, list[str]] = {}
    failed: list[tuple[int, str]] = []
    for slide, text, reason in results:
        text = text.strip()
        if reason:
            failed.append((slide, reason))
        elif text:
            by_slide.setdefault(slide, []).append(text)
    pages, added = [], {}
    for number, body in reading.pages:
        found_texts = by_slide.get(number)
        if found_texts:
            addition = '\n\n'.join(f'[图片内容]\n{text}' for text in found_texts)
            added[number] = addition
            body = (body + '\n\n' if body.strip() else '') + addition
        pages.append((number, body))

    notes = _drop_old_notes(reading.notes)
    if by_slide:
        notes.append(f'第 {_numbers(by_slide)} 页的图片由识图模型识别，可能有错字、漏字，示意图和公式请对照原件。')
    if failed:
        notes.append(f'第 {_numbers(slide for slide, _ in failed)} 页有图片识别失败（{failed[0][1][:60]}），那些图片没有读取。')
    if found.over_limit:
        notes.append(f'课件图片很多，只识别了前 {MAX_PICTURES} 张，其余 {found.over_limit} 张没有读取。')
    if found.unsupported:
        notes.append(f'有 {found.unsupported} 张矢量图（EMF/WMF）没有识别。')
    if not by_slide and not failed:
        notes.append('课件里的图片没有识别到文字或有效信息。')
    return Reading(pages, notes, extra={**reading.extra, **added})
